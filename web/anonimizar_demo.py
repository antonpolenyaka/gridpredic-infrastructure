"""
Anonimiza la copia de datos de la demo pública (web/demo/data) antes de publicarla. Los JSON y los partes de
guardia que dejan los jobs llevan los nombres reales de los municipios de cada red (en los municipios, las zonas,
los CT, las bases de las brigadas, el pico de AEMET y los textos), su código INE y las estaciones de AEMET. Con
ellos se puede saber qué distribuidora es cada una, y la memoria no los da. Este guion los cambia por seudónimos:

- cada municipio pasa a "Municipio N" en todos los ficheros (las zonas, que se llaman por sus municipios, quedan
  como "Municipio N / Municipio M", y el cuadro de mando sigue cruzándolas con sus municipios);
- su código INE, a 99NNN, que no es el de ningún municipio;
- cada estación de AEMET, a "Estación N", y su indicativo, a ENN.

Los números se sortean la primera vez y se guardan en el mapa (por defecto _runlogs/demo_seudonimos.json, que no
se versiona porque lleva los nombres reales); las copias siguientes usan el mismo mapa, así que un municipio tiene
el mismo seudónimo en todos los días archivados. sector.json no se toca: compara con la provincia, la comunidad y
España, como la memoria, y no lleva municipios. Al final busca en todo lo escrito los nombres, los códigos y los
indicativos reales (sin distinguir mayúsculas ni tildes) y para con un error si queda alguno.

    python web/anonimizar_demo.py web/demo/data
    python web/anonimizar_demo.py web/demo/data --mapa _runlogs/demo_seudonimos.json

Necesita openpyxl (está en requirements.txt).
"""

import argparse
import json
import os
import random
import re
import sys
import unicodedata

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DEFAULT_MAP = os.path.join(ROOT, "_runlogs", "demo_seudonimos.json")
SKIP = {"sector.json"}
PSEUDO_TOWN = re.compile(r"^Municipio \d+$")
PSEUDO_STATION = re.compile(r"^Estación \d+$")
INE_HEADER = "código INE"


def fold(text: str) -> str:
    """Lower case and without accents, to look for what is left."""
    return "".join(c for c in unicodedata.normalize("NFD", text.lower()) if unicodedata.category(c) != "Mn")


def read_json(path: str):
    with open(path, encoding="utf-8") as handle:
        return json.load(handle)


def layout(raw: str) -> dict:
    """How the job wrote the file (indented with n spaces or compact), so that only the names change in the diff."""
    m = re.match(r"\{\r?\n( +)\"", raw)
    if m:
        return {"indent": len(m.group(1))}
    return {"separators": (",", ":")}


def write_json(path: str, data, raw: str) -> None:
    """UTF-8 without escaping and with the layout of the original; to a temporary file and renamed."""
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="") as handle:
        json.dump(data, handle, ensure_ascii=False, **layout(raw))
        if raw.endswith("\n"):
            handle.write("\n")
    os.replace(tmp, path)


def collect(folder: str) -> tuple:
    """({municipio: código INE}, {estación: indicativo}) reales de la carpeta, sin los que ya son seudónimos."""
    towns, stations = {}, {}
    for name in sorted(os.listdir(folder)):
        path = os.path.join(folder, name)
        if name.startswith("riesgo") and name.endswith(".json"):
            for dist in read_json(path).get("distribuidoras", []):
                for town in dist.get("municipios", []):
                    if town.get("nombre") and not PSEUDO_TOWN.match(town["nombre"]):
                        towns[town["nombre"]] = str(town.get("codigo_ine") or "")
        elif name.startswith("observado") and name.endswith(".json"):
            for station in read_json(path).get("estaciones", []):
                if station.get("nombre") and not PSEUDO_STATION.match(station["nombre"]):
                    stations[station["nombre"]] = str(station.get("idema") or "")
    return towns, stations


def extend(mapping: dict, towns: dict, stations: dict) -> int:
    """Adds the new names to the map with numbers drawn at random after the ones already used. Returns how many."""
    rng = random.SystemRandom()
    added = 0
    for kind, found, label in (("municipios", towns, "Municipio"), ("estaciones", stations, "Estación")):
        table = mapping.setdefault(kind, {})
        new = [n for n in found if n not in table]
        rng.shuffle(new)
        used = [int(v["seudonimo"].rsplit(" ", 1)[1]) for v in table.values()]
        for i, real in enumerate(new, start=max(used, default=0) + 1):
            entry = {"seudonimo": f"{label} {i}"}
            if kind == "municipios":
                entry.update(codigo_real=found[real], codigo=f"99{i:03d}")
            else:
                entry.update(indicativo_real=found[real], indicativo=f"E{i:02d}")
            table[real] = entry
            added += 1
        for real, code in found.items():
            # A code that was missing the first time.
            key = "codigo_real" if kind == "municipios" else "indicativo_real"
            if code and not table[real].get(key):
                table[real][key] = code
    return added


class Replacer:
    def __init__(self, mapping: dict):
        towns, stations = mapping.get("municipios", {}), mapping.get("estaciones", {})
        self.names = {real: e["seudonimo"] for real, e in {**towns, **stations}.items()}
        # Longest first, so "Valle de Santa Ana" goes before "Santa Ana" and "Casar de Cáceres" before "Cáceres".
        alternatives = sorted(self.names, key=len, reverse=True)
        self.pattern = re.compile(r"(?<!\w)(" + "|".join(map(re.escape, alternatives)) + r")(?!\w)") if alternatives else None
        self.codes = {e["codigo_real"]: e["codigo"] for e in towns.values() if e.get("codigo_real")}
        self.ids = {e["indicativo_real"]: e["indicativo"] for e in stations.values() if e.get("indicativo_real")}

    def text(self, value: str) -> str:
        return self.pattern.sub(lambda m: self.names[m.group(1)], value) if self.pattern else value

    def data(self, value, key: str = ""):
        if isinstance(value, dict):
            return {self.text(k): self.data(v, k) for k, v in value.items()}
        if isinstance(value, list):
            return [self.data(v, key) for v in value]
        if isinstance(value, str):
            if key == "codigo_ine":
                return self.codes.get(value, value)
            if key == "idema":
                return self.ids.get(value, value)
            return self.text(value)
        return value


def anonymize_xlsx(path: str, rep: Replacer) -> bool:
    from openpyxl import load_workbook

    book = load_workbook(path)
    changed = False
    for sheet in book.worksheets:
        code_cols = set()
        for row in sheet.iter_rows():
            for cell in row:
                old = cell.value
                if cell.row == 1 and old == INE_HEADER:
                    code_cols.add(cell.column)
                elif cell.row > 1 and cell.column in code_cols and old is not None:
                    cell.value = rep.codes.get(str(old), old)
                elif isinstance(old, str):
                    cell.value = rep.text(old)
                changed = changed or cell.value != old
    if not changed:
        return False
    tmp = path + ".tmp.xlsx"
    book.save(tmp)
    os.replace(tmp, path)
    return True


def xlsx_contents(path: str) -> tuple:
    """(all the text of the workbook, the values of its INE code columns)."""
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True)
    parts, codes = [], []
    for sheet in book.worksheets:
        col = None
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i == 0 and INE_HEADER in row:
                col = row.index(INE_HEADER)
            parts.extend(str(v) for v in row if v is not None)
            if i > 0 and col is not None and len(row) > col and row[col] is not None:
                codes.append(str(row[col]))
    book.close()
    return "\n".join(parts), codes


def json_fields(data, wanted: set) -> list:
    """(key, value) of every field of the JSON whose key is in wanted."""
    out, stack = [], [("", data)]
    while stack:
        key, value = stack.pop()
        if isinstance(value, dict):
            stack.extend(value.items())
        elif isinstance(value, list):
            stack.extend((key, v) for v in value)
        elif key in wanted:
            out.append((key, str(value)))
    return out


def leftovers(folder: str, mapping: dict) -> list:
    """Real names, INE codes and station ids still in the files of the folder (sector.json aside)."""
    towns, stations = mapping.get("municipios", {}), mapping.get("estaciones", {})
    words = sorted({fold(n) for n in list(towns) + list(stations)}, key=len, reverse=True)
    names = re.compile(r"(?<!\w)(" + "|".join(map(re.escape, words)) + r")(?!\w)") if words else None
    codes = {e["codigo_real"] for e in towns.values() if e.get("codigo_real")}
    ids = {e["indicativo_real"] for e in stations.values() if e.get("indicativo_real")}
    found = []
    for name in sorted(os.listdir(folder)):
        path = os.path.join(folder, name)
        if name in SKIP or not name.endswith((".json", ".xlsx")):
            continue
        if name.endswith(".xlsx"):
            text, in_column = xlsx_contents(path)
            fields = [("codigo_ine", c) for c in in_column]
        else:
            with open(path, encoding="utf-8") as handle:
                text = handle.read()
            fields = json_fields(json.loads(text), {"codigo_ine", "idema"})
        hit = names.search(fold(text)) if names else None
        if hit:
            found.append(f"{name}: el nombre real {hit.group(1)!r}")
        bad = [v for k, v in fields if (k == "codigo_ine" and v in codes) or (k == "idema" and v in ids)]
        if bad:
            found.append(f"{name}: {len(bad)} códigos o indicativos reales")
    return found


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Seudónimos en la copia de datos de la demo pública")
    parser.add_argument("carpeta", help="web/demo/data o la carpeta que se va a publicar")
    parser.add_argument("--mapa", default=DEFAULT_MAP, help="mapa de seudónimos (no se versiona)")
    args = parser.parse_args(argv)

    mapping = read_json(args.mapa) if os.path.exists(args.mapa) else {}
    towns, stations = collect(args.carpeta)
    added = extend(mapping, towns, stations)
    rep = Replacer(mapping)

    changed = 0
    for name in sorted(os.listdir(args.carpeta)):
        path = os.path.join(args.carpeta, name)
        if name in SKIP:
            continue
        if name.endswith(".json"):
            with open(path, encoding="utf-8") as handle:
                raw = handle.read()
            data = json.loads(raw)
            new = rep.data(data)
            if new != data:
                write_json(path, new, raw)
                changed += 1
        elif name.endswith(".xlsx") and anonymize_xlsx(path, rep):
            changed += 1

    os.makedirs(os.path.dirname(os.path.abspath(args.mapa)), exist_ok=True)
    tmp = args.mapa + ".tmp"
    with open(tmp, "w", encoding="utf-8") as handle:
        json.dump(mapping, handle, ensure_ascii=False, indent=1, sort_keys=True)
    os.replace(tmp, args.mapa)

    left = leftovers(args.carpeta, mapping)
    print(f"{len(mapping.get('municipios', {}))} municipios y {len(mapping.get('estaciones', {}))} estaciones en el mapa "
          f"({added} nuevos); {changed} ficheros escritos en {args.carpeta}")
    if left:
        print("ERROR: quedan datos reales:\n  " + "\n  ".join(left), file=sys.stderr)
        return 1
    print("Sin nombres, códigos INE ni indicativos reales.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
