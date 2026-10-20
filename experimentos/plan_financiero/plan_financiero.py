# -*- coding: utf-8 -*-
"""Plan financiero de GridPredic a tres años (capítulo 3 de la memoria, 3.9).

Modelo mensual con partida doble: cada movimiento de caja tiene su contrapartida y el
balance se calcula desde las cuentas, sin partida de ajuste. En cada ejecución
comprobamos que el activo es igual al patrimonio neto más el pasivo en los 37 cierres
(momento 0 y meses 1 a 36), que la tesorería final de un año es la inicial del
siguiente y que el resultado de cada año pasa a los fondos propios. Si algo no
cuadra, el guion se para con un AssertionError.

Hay dos juegos de supuestos:

- V10: el plan de la memoria v10 (licencia de 18.000 euros, ENISA de 38.000 euros
  cobrada en el momento 0, 40.000 euros de los fundadores). Lo reproducimos al euro
  como referencia.
- V11: las decisiones del 19 y del 20.10.2026. Licencia por tamaño de red, con el
  precio por CT en tramos (descuento por volumen), un mínimo y sin máximo; ENISA de
  35.000 euros cobrada en el mes 6; 40.000 euros de los fundadores con 3.000 de
  capital, y medio sueldo de los fundadores hasta que llega ENISA. Con eso la caja no
  baja de cero y el patrimonio neto no baja de la mitad del capital (art. 363.1.e LSC)
  en ningún mes, y el guion lo comprueba.

Para comparar los dos planes más allá del año 3 alargamos los dos a 60 meses con los
mismos supuestos de los años 4 y 5 (ANIOS_4_5). Esa extensión es solo de contraste:
las tablas de 3.9 son las de 36 meses.

Uso:
    python plan_financiero.py              # calcula, comprueba y escribe salida/
    python plan_financiero.py --sin-salida # solo calcula, comprueba y enseña las cifras

Salidas (salida/):
    tablas_v11.txt       tablas de 3.9 en el formato de la memoria (T, H, W, R)
    tablas_v10.txt       las mismas tablas con los supuestos de la v10
    cifras_v11.json      cifras sueltas que cita el texto
    comparacion.txt      antes y después de las cifras clave, con la licencia de cada red
    mensual_v11.csv      tesorería, resultado y patrimonio de cada mes
    detalle_v11.xlsx     el detalle mes a mes y los balances (si está openpyxl)

Todas las cifras van sin IVA salvo los cobros y pagos de la tesorería. Solo usa la
biblioteca estándar (openpyxl es opcional).
"""

import copy
import csv
import json
import math
import os
import sys

AQUI = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.join(AQUI, "salida")

# ---------------------------------------------------------------------------
# Tarifa por tamaño de red (decisión 9 del 19.10.2026; forma elegida el 20.10.2026)
# ---------------------------------------------------------------------------
# Como se licencia un SCADA: una parte fija, una por CT y otra por señal (los tags del
# SCADA que cuelgan de la distribuidora). Cada CT paga el precio de su tramo, como en el
# IRPF, así que la licencia no da saltos al pasar de un tramo a otro. El precio por CT
# baja a partir de 100 CT porque el coste de servir a una red apenas crece con los CT:
# el modelo, la previsión y el cuadro de mando son los mismos. Mínimo para las redes muy
# pequeñas y sin máximo.
TARIFA = {
    "fijo": 3000.0,
    "tramos_ct": [(100, 150.0), (300, 15.0), (None, 10.0)],   # (hasta el CT, euros por CT)
    "por_senal": 0.10,
    "minimo": 6000.0,
}

# ÚNICO SITIO con los CT y las señales de las tres redes del TFM. Salen de
# senales_por_distribuidora.sql sobre la copia de TedisNet de las tres redes (20.10.2026):
# CT = elementos de tipo TRAFO CT de cada distribuidora (pueden diferir algo de los CT de
# Calser, unos 760, 570 y 10); señales = tags cuyo elemento cuelga de la distribuidora. Los
# 14.705 tags sin elemento no se reparten. Si cambian, se vuelve a ejecutar el guion y se
# copian las licencias nuevas de salida/comparacion.txt.
REDES_CALIBRACION = {
    "Distribuidora 1": {"ct": 796, "senales": 34647},
    "Distribuidora 2": {"ct": 578, "senales": 27513},
    "Distribuidora 3": {"ct": 10, "senales": 606},
}
SENALES_PROVISIONALES = False

# Red tipo de los clientes del plan: una red pequeña con telecontrol, cerca del centro de
# las 89 distribuidoras más grandes (39 de ellas tienen 150 CT o más y 62 tienen 100 CT o más,
# estimados por su retribución). Su licencia es la de cada cliente del plan.
RED_TIPO_PLAN = {"ct": 150, "senales": 9000}

# Redes de referencia para la página de precios y para la sensibilidad
REDES_REFERENCIA = {"Red de 50 CT": {"ct": 50, "senales": 2900},
                    "Red de 100 CT": {"ct": 100, "senales": 6000},
                    "Red de 300 CT": {"ct": 300, "senales": 17500}}


def partes_licencia(ct, senales, t=TARIFA):
    """Fijo, parte de los CT (por tramos) y parte de las señales, antes del mínimo."""
    parte_ct = 0.0
    desde = 0
    for hasta, precio in t["tramos_ct"]:
        tope = ct if hasta is None else min(ct, hasta)
        if tope > desde:
            parte_ct += (tope - desde) * precio
        if hasta is None or ct <= hasta:
            break
        desde = hasta
    return t["fijo"], parte_ct, t["por_senal"] * senales


def licencia(ct, senales, t=TARIFA):
    """Licencia anual en euros: fijo + CT por tramos + señales, con el mínimo y sin techo."""
    return max(t["minimo"], sum(partes_licencia(ct, senales, t)))


def texto_tarifa(t=TARIFA, moneda="€"):
    (h1, p1), (h2, p2), (_, p3) = t["tramos_ct"]
    return (f"{fmt(t['fijo'])} {moneda} fijos, {fmt(p1)} {moneda} por CT del 1 al {h1}, {fmt(p2)} {moneda} del "
            f"{h1 + 1} al {h2} y {fmt(p3)} {moneda} desde el {h2 + 1}, y {fmt2(t['por_senal'])} {moneda} por señal, "
            f"con un mínimo de {fmt(t['minimo'])} {moneda} y sin máximo")


# ---------------------------------------------------------------------------
# Supuestos
# ---------------------------------------------------------------------------
V10 = {
    "nombre": "v10",
    "meses": 36,
    # mes de puesta en marcha (1 = enero de 2027); la firma es el mes anterior
    "puestas_en_marcha": [5, 7, 9,
                          14, 16, 18, 20, 22, 24,
                          26, 27, 29, 30, 32, 33, 35, 36],
    # (mes de puesta en marcha del cliente, aniversario en que no renueva)
    "bajas": [(9, 21), (14, 26)],
    "precio_implantacion": 20000.0,
    "precio_licencia": 18000.0,
    "parte_implantacion_firma": 0.5,
    "cobro_meses": 1,
    "plantilla": [("Fundador (dirección y negocio)", 50000.0, 1),
                  ("Fundador (datos y plataforma)", 40000.0, 1),
                  ("Ingeniero de implantación y soporte", 30000.0, 13),
                  ("Responsable comercial del sector", 36000.0, 25)],
    # (desde, hasta, parte del sueldo de los fundadores a la que renuncian); si dos
    # periodos se solapan, cuenta la rebaja mayor
    "recortes_fundadores": [],
    "cargas_sociales": 0.30,
    "ss_trabajador": 0.065,
    "retencion_irpf": 0.18,
    # euros al año; el último valor se repite si el horizonte es más largo
    "gastos_anuales": {
        "Nube (cómputo, almacenamiento, demos)": [7200.0, 9600.0, 12000.0],
        "Asesoría jurídica": [4000.0, 8000.0, 6000.0],
        "Gestoría": [2400.0, 3600.0, 4800.0],
        "Actividad comercial y ferias": [5000.0, 15000.0, 25000.0],
        "Equipos menores, viajes y otros": [3000.0, 6000.0, 9000.0],
        "Seguros (responsabilidad civil)": [1200.0, 2000.0, 2000.0],
    },
    "herramientas_persona_mes": 100.0,
    "coworking_persona_mes": 250.0,
    "comision_socio": 0.10,
    "comision_desde_mes": 13,
    "imprevistos": 0.05,
    "inversiones": [(13, "Equipo del ingeniero", 2500.0, "equipo"),
                    (14, "Marca en la EUIPO", 1000.0, "marca"),
                    (25, "Equipo del responsable comercial", 2500.0, "equipo")],
    "amortizacion": {"marca": 0.10, "aplicaciones": 0.25, "mobiliario": 0.10, "equipo": 0.25},
    "reparto_equipo": {"equipo": 2000.0 / 2500.0, "mobiliario": 500.0 / 2500.0},
    "activo_inicial": {"marca": 500.0, "aplicaciones": 5000.0, "mobiliario": 1500.0, "equipo": 6000.0},
    "fianza": 1000.0,
    "capital_fundadores": 3000.0,
    "prima_fundadores": 37000.0,
    "gastos_constitucion": 1500.0,
    # mes_cobro 0 = en el momento 0, como la v10
    "enisa": {"importe": 38000.0, "interes": 0.06, "carencia_meses": 24, "plazo_meses": 84,
              "comision_apertura": 0.005, "mes_cobro": 0},
    "banco": {"importe": 20000.0, "interes": 0.065, "carencia_meses": 12, "plazo_meses": 60},
    "ronda": {"mes": 12, "importe": 250000.0, "participacion": 0.25, "gastos": 3000.0},
    "iva": 0.21,
    "tipo_is": 0.15,
    "mes_pago_is": 7,
    # subida anual de la licencia en cada aniversario (el plan no la cuenta; ver valor_ipc)
    "ipc_licencia": 0.0,
    # {mes previsto de la contratación: puesta en marcha con la que entra}; vacío = sin regla
    "contrataciones_ligadas": {},
}

# ENISA (decisión 10). Constituimos en diciembre de 2026 y el primer ejercicio cierra el
# 31.12.2026. En enero formulamos, aprobamos en junta universal y depositamos las cuentas
# de 2026 (arts. 253.1, 164.1, 178 y 279.1 LSC); la solicitud va en febrero (mes 2); ENISA
# da 30 días para completar el expediente y no publica un plazo de resolución (las
# asesorías hablan de uno a tres meses con el expediente completo); con la póliza y el
# desembolso, el dinero está en el mes 6. Pedimos 35.000 euros y no 38.000 para tener
# margen en la condición de fondos propios (al menos iguales al préstamo), que se mide
# con las cuentas de 2026: 38.500 euros.
MES_ENISA_V11 = 6
IMPORTE_ENISA_V11 = 35000.0

# Capital y aportaciones (decisión 11, sin Neotec). El capital se queda en 3.000 euros:
# con menos se aplica el régimen de las sociedades limitadas de menos de 3.000 euros
# (art. 4.1 LSC, redacción de la Ley 18/2022) y con más sube el umbral de disolución, la
# mitad del capital (art. 363.1.e). Los fundadores ponemos 40.000 euros (3.000 de capital y
# 37.000 de prima de asunción) y cobramos medio sueldo de enero a mayo de 2027 (meses 1 a
# 5), hasta que llega ENISA en junio: es una
# rebaja de verdad, no un aplazamiento, porque un sueldo aplazado se devenga igual y no
# mejora el patrimonio. calcular() comprueba que la rebaja basta.
APORTACION_FUNDADORES_V11 = 40000.0
RECORTE_FUNDADORES_V11 = (1, MES_ENISA_V11 - 1, 0.5)

V11 = copy.deepcopy(V10)
V11.update({
    "nombre": "v11",
    "precio_licencia": licencia(**RED_TIPO_PLAN),
    "prima_fundadores": APORTACION_FUNDADORES_V11 - V10["capital_fundadores"],
    "recortes_fundadores": [RECORTE_FUNDADORES_V11],
    # El ingeniero entra con la tercera puesta en marcha y el comercial con la novena, nunca
    # antes de los meses 13 y 25. En el plan no cambia nada; en los retrasos, sí.
    "contrataciones_ligadas": {13: 3, 25: 9},
})
V11["enisa"]["mes_cobro"] = MES_ENISA_V11
V11["enisa"]["importe"] = IMPORTE_ENISA_V11

# Años 4 y 5, solo para comparar los dos planes (los mismos para los dos): ocho clientes
# nuevos al año al ritmo del año 3, dos bajas al año, la misma plantilla y los gastos del
# año 3.
ANIOS_4_5 = {
    "puestas_extra": [g + 12 for g in (26, 27, 29, 30, 32, 33, 35, 36)]
    + [g + 24 for g in (26, 27, 29, 30, 32, 33, 35, 36)],
    "bajas_extra": [(16, 40), (18, 42), (27, 51), (29, 53)],
}
TASA_VAN = 0.10   # tasa de contraste; la misma para los dos planes

# Hipótesis de venta para la cuenta del inversor (3.9.6): un múltiplo de las licencias
# anuales y el plazo desde la entrada de la ronda (mes 12) hasta la venta a finales de 2031.
MULTIPLO_VENTA = 4.0
ANIOS_HASTA_VENTA = 4
SIN_IVA = {"Seguros (responsabilidad civil)"}


def a_60_meses(S):
    S2 = copy.deepcopy(S)
    S2["meses"] = 60
    S2["puestas_en_marcha"] = S["puestas_en_marcha"] + ANIOS_4_5["puestas_extra"]
    S2["bajas"] = S["bajas"] + ANIOS_4_5["bajas_extra"]
    return S2


# ---------------------------------------------------------------------------
# Formato
# ---------------------------------------------------------------------------
def r0(x):
    return int(round(x + (1e-9 if x >= 0 else -1e-9)))


def fmt(x, cero="0"):
    v = r0(x)
    if v == 0:
        return cero
    s = f"{abs(v):,}".replace(",", ".")
    return ("-" if v < 0 else "") + s


def fmt1(x):
    """Un decimal con coma: 18,6."""
    return f"{x:.1f}".replace(".", ",")


def fmt2(x):
    """Dos decimales con coma: 0,10."""
    return f"{x:.2f}".replace(".", ",")


def fmt_pct(x, dec=1):
    return f"{x * 100:.{dec}f}".replace(".", ",") + " %"


def anio_de(m):
    return (m - 1) // 12 + 1


# ---------------------------------------------------------------------------
# Simulación
# ---------------------------------------------------------------------------
def meses_contratacion(S):
    """Mes en que entra cada contratación con la regla de las puestas en marcha:
    {mes previsto: mes real}; None si no llegan las puestas en marcha que hacen falta."""
    pm = sorted(S["puestas_en_marcha"])
    return {mes: (max(mes, pm[n - 1]) if len(pm) >= n else None)
            for mes, n in S.get("contrataciones_ligadas", {}).items()}


def aplicar_contrataciones(S):
    meses = meses_contratacion(S)
    if not meses:
        return S
    S2 = copy.deepcopy(S)
    nunca = 10 ** 6
    S2["plantilla"] = [(n, b, (meses[d] or nunca) if d in meses else d) for (n, b, d) in S["plantilla"]]
    S2["inversiones"] = [(((meses[m] or nunca) if (m in meses and tipo == "equipo") else m), c, imp, tipo)
                         for (m, c, imp, tipo) in S["inversiones"]]
    return S2


def precio_licencia_anio(S, k):
    """Licencia del año k del cliente (0 = el primero), con la subida anual si la hay."""
    return S["precio_licencia"] * (1 + S.get("ipc_licencia", 0.0)) ** k


def simular(S):
    S = aplicar_contrataciones(copy.deepcopy(S))
    N = S["meses"]
    iva = S["iva"]
    P = {k: [0.0] * (N + 2) for k in [
        "ing_impl", "ing_lic", "sueldos", "cargas", "externos", "comision", "imprevistos",
        "amort", "financieros", "is",
        "fact_base", "fact_iva", "cobros", "pago_nomina", "pago_ss", "pago_irpf",
        "pago_externos", "pago_iva", "pago_is", "pago_inversion", "pago_intereses",
        "pago_principal", "entrada_capital", "pago_gastos_ronda", "entrada_prestamo",
    ]}
    detalle_ext = {}

    # --- clientes -------------------------------------------------------------
    bajas = dict(S["bajas"])
    clientes = [{"gm": gm, "fin": bajas.get(gm)} for gm in S["puestas_en_marcha"]]
    for c in clientes:
        gm = c["gm"]
        firma = gm - 1
        f1 = S["precio_implantacion"] * S["parte_implantacion_firma"]
        f2 = S["precio_implantacion"] - f1
        if 1 <= firma <= N:
            P["fact_base"][firma] += f1
        if gm <= N:
            P["fact_base"][gm] += f2
            P["ing_impl"][gm] += S["precio_implantacion"]
        m = gm
        while m <= N:
            if c["fin"] is not None and m >= c["fin"]:
                break
            P["fact_base"][m] += precio_licencia_anio(S, (m - gm) // 12)
            m += 12
        for m in range(gm, N + 1):
            if c["fin"] is not None and m >= c["fin"]:
                break
            P["ing_lic"][m] += precio_licencia_anio(S, (m - gm) // 12) / 12.0
    for m in range(1, N + 1):
        P["fact_iva"][m] = P["fact_base"][m] * iva

    # --- plantilla --------------------------------------------------------------
    personas_mes = [0] * (N + 2)
    for (_, bruto, desde) in S["plantilla"]:
        for m in range(desde, N + 1):
            P["sueldos"][m] += bruto / 12.0
            personas_mes[m] += 1
    # rebaja del sueldo de los fundadores (renuncia, no aplazamiento); si dos periodos se
    # solapan, cuenta la rebaja mayor
    parte_mes = [0.0] * (N + 2)
    for (d, h, parte) in S.get("recortes_fundadores", []):
        for m in range(d, min(h, N) + 1):
            parte_mes[m] = max(parte_mes[m], parte)
    bruto_fund = sum(b for (_, b, desde) in S["plantilla"] if desde == 1)
    for m in range(1, N + 1):
        P["sueldos"][m] -= bruto_fund / 12.0 * parte_mes[m]
    for m in range(1, N + 1):
        P["cargas"][m] = P["sueldos"][m] * S["cargas_sociales"]

    # --- gastos externos ----------------------------------------------------------
    for concepto, anual in S["gastos_anuales"].items():
        detalle_ext[concepto] = [0.0] * (N + 2)
        for m in range(1, N + 1):
            detalle_ext[concepto][m] = anual[min(anio_de(m), len(anual)) - 1] / 12.0
    detalle_ext["Herramientas de desarrollo"] = [0.0] * (N + 2)
    detalle_ext["Coworking"] = [0.0] * (N + 2)
    for m in range(1, N + 1):
        detalle_ext["Herramientas de desarrollo"][m] = S["herramientas_persona_mes"] * personas_mes[m]
        detalle_ext["Coworking"][m] = S["coworking_persona_mes"] * personas_mes[m]
    for m in range(1, N + 1):
        P["externos"][m] = sum(v[m] for v in detalle_ext.values())
        if m >= S["comision_desde_mes"]:
            P["comision"][m] = P["ing_lic"][m] * S["comision_socio"]
        base_imp = P["sueldos"][m] + P["cargas"][m] + P["externos"][m] + P["comision"][m]
        P["imprevistos"][m] = base_imp * S["imprevistos"]

    iva_sop = [0.0] * (N + 2)
    for m in range(1, N + 1):
        con_iva = sum(v[m] for k, v in detalle_ext.items() if k not in SIN_IVA)
        con_iva += P["comision"][m] + P["imprevistos"][m]
        iva_sop[m] = con_iva * iva
        P["pago_externos"][m] = P["externos"][m] + P["comision"][m] + P["imprevistos"][m] + iva_sop[m]

    # --- inmovilizado -------------------------------------------------------------
    activos = [(1, tipo, coste) for tipo, coste in S["activo_inicial"].items()]
    for (mes, _, importe, tipo) in S["inversiones"]:
        if mes > N:
            continue
        if tipo == "equipo":
            for t2, parte in S["reparto_equipo"].items():
                activos.append((mes, t2, importe * parte))
        else:
            activos.append((mes, tipo, importe))
        P["pago_inversion"][mes] += importe * (1 + iva)
        iva_sop[mes] += importe * iva
    amort_acum = {"intangible": [0.0] * (N + 2), "material": [0.0] * (N + 2)}
    coste_bruto = {"intangible": [0.0] * (N + 2), "material": [0.0] * (N + 2)}
    for m in range(1, N + 1):
        for (alta, tipo, coste) in activos:
            clase = "intangible" if tipo in ("marca", "aplicaciones") else "material"
            if alta <= m:
                coste_bruto[clase][m] += coste
                cuota = coste * S["amortizacion"][tipo] / 12.0
                acum_hoy = min(cuota * (m - alta + 1), coste)
                acum_antes = min(cuota * (m - alta), coste)
                amort_acum[clase][m] += acum_hoy
                P["amort"][m] += acum_hoy - acum_antes

    # --- financiación -------------------------------------------------------------
    E, B, R = S["enisa"], S["banco"], S["ronda"]
    mc = E["mes_cobro"]
    enisa_pend = [0.0] * (N + 2)
    banco_pend = [0.0] * (N + 2)
    banco_pend[0] = B["importe"]
    if mc == 0:
        enisa_pend[0] = E["importe"]
    cuota_enisa = E["importe"] / ((E["plazo_meses"] - E["carencia_meses"]) / 3.0)   # trimestral
    cuota_banco = B["importe"] / (B["plazo_meses"] - B["carencia_meses"])           # mensual
    comision = E["importe"] * E["comision_apertura"]
    P["financieros"][max(mc, 1)] += comision
    P["pago_intereses"][max(mc, 1)] += comision
    intereses_pend = [0.0] * (N + 2)
    for m in range(1, N + 1):
        e_prev = enisa_pend[m - 1]
        b_prev = banco_pend[m - 1]
        if mc > 0 and m == mc:
            P["entrada_prestamo"][m] += E["importe"]
        P["financieros"][m] += e_prev * E["interes"] / 12.0
        if m > mc and (m - mc) % 3 == 0:
            P["pago_intereses"][m] += sum(enisa_pend[k] * E["interes"] / 12.0 for k in range(m - 3, m))
            if m > mc + E["carencia_meses"]:
                pago = min(cuota_enisa, e_prev)
                P["pago_principal"][m] += pago
                e_prev -= pago
        if mc > 0 and m == mc:
            e_prev += E["importe"]
        enisa_pend[m] = e_prev
        P["financieros"][m] += b_prev * B["interes"] / 12.0
        P["pago_intereses"][m] += b_prev * B["interes"] / 12.0
        if m > B["carencia_meses"]:
            pago = min(cuota_banco, b_prev)
            P["pago_principal"][m] += pago
            b_prev -= pago
        banco_pend[m] = b_prev
        intereses_pend[m] = intereses_pend[m - 1] + P["financieros"][m] - P["pago_intereses"][m]
    # la comisión se paga el mismo mes que se devenga: no queda pendiente
    capital_nominal_fund = S["capital_fundadores"]
    nuevas = 0
    if R["mes"] and R["mes"] <= N:
        nuevas = int(-(-capital_nominal_fund * R["participacion"] // (1 - R["participacion"])))
        P["entrada_capital"][R["mes"]] += R["importe"]
        P["pago_gastos_ronda"][R["mes"]] += R["gastos"] * (1 + iva)
        iva_sop[R["mes"]] += R["gastos"] * iva

    # --- impuestos y tesorería ------------------------------------------------------
    caja = [0.0] * (N + 2)
    caja[0] = (S["capital_fundadores"] + S["prima_fundadores"] + enisa_pend[0] + B["importe"]
               - S["gastos_constitucion"] - sum(S["activo_inicial"].values()) - S["fianza"])
    iva_saldo = [0.0] * (N + 2)
    irpf_pend = [0.0] * (N + 2)
    ss_pend = [0.0] * (N + 2)
    is_pend = [0.0] * (N + 2)
    bin_disponible = 0.0
    resultados_anuales, bai_anual, is_anual, iva_fin_trimestre = {}, {}, {}, {}
    for m in range(1, N + 1):
        k = m - S["cobro_meses"]
        if k >= 1:
            P["cobros"][m] = P["fact_base"][k] + P["fact_iva"][k]
        bruto = P["sueldos"][m]
        P["pago_nomina"][m] = bruto * (1 - S["retencion_irpf"] - S["ss_trabajador"])
        P["pago_ss"][m] = ss_pend[m - 1]
        ss_pend[m] = bruto * (S["cargas_sociales"] + S["ss_trabajador"])
        irpf_pend[m] = irpf_pend[m - 1] + bruto * S["retencion_irpf"]
        if m % 3 == 1 and m > 1:
            P["pago_irpf"][m] = irpf_pend[m - 1]
            irpf_pend[m] -= irpf_pend[m - 1]
        iva_saldo[m] = iva_saldo[m - 1] + P["fact_iva"][m] - iva_sop[m]
        if m % 3 == 0:
            iva_fin_trimestre[m] = iva_saldo[m]
        if m % 3 == 1 and m > 1:
            a_pagar = iva_fin_trimestre[m - 1]
            if a_pagar > 0:
                P["pago_iva"][m] = a_pagar
                iva_saldo[m] -= a_pagar
        if (m - 1) % 12 + 1 == S["mes_pago_is"] and m > 12:
            P["pago_is"][m] = is_pend[m - 1]
            is_pend[m] = is_pend[m - 1] - P["pago_is"][m]
        else:
            is_pend[m] = is_pend[m - 1]
        if m % 12 == 0:
            a = m // 12
            meses_a = range(m - 11, m + 1)
            ingresos = sum(P["ing_impl"][x] + P["ing_lic"][x] for x in meses_a)
            gastos = sum(P["sueldos"][x] + P["cargas"][x] + P["externos"][x] + P["comision"][x]
                         + P["imprevistos"][x] + P["amort"][x] + P["financieros"][x] for x in meses_a)
            bai = ingresos - gastos
            if bai < 0:
                bin_disponible += -bai
                cuota = 0.0
            else:
                compensa = min(bin_disponible, bai)
                bin_disponible -= compensa
                cuota = (bai - compensa) * S["tipo_is"]
            P["is"][m] = cuota
            is_pend[m] += cuota
            bai_anual[a], is_anual[a], resultados_anuales[a] = bai, cuota, bai - cuota
        salidas = (P["pago_nomina"][m] + P["pago_ss"][m] + P["pago_irpf"][m] + P["pago_externos"][m]
                   + P["pago_iva"][m] + P["pago_is"][m] + P["pago_inversion"][m] + P["pago_intereses"][m]
                   + P["pago_principal"][m] + P["pago_gastos_ronda"][m])
        entradas = P["cobros"][m] + P["entrada_capital"][m] + P["entrada_prestamo"][m]
        caja[m] = caja[m - 1] + entradas - salidas

    # --- balances -------------------------------------------------------------------
    def corto_plazo(m):
        e_corto = sum(cuota_enisa for k in range(m + 1, m + 13)
                      if k > mc and (k - mc) % 3 == 0 and k > mc + E["carencia_meses"])
        e_corto = min(e_corto, enisa_pend[m])
        b_corto = min(sum(cuota_banco for k in range(m + 1, m + 13) if k > B["carencia_meses"]), banco_pend[m])
        return e_corto, b_corto

    def balance(m):
        b = {}
        if m == 0:
            b["intangible"] = S["activo_inicial"]["marca"] + S["activo_inicial"]["aplicaciones"]
            b["material"] = S["activo_inicial"]["mobiliario"] + S["activo_inicial"]["equipo"]
        else:
            b["intangible"] = coste_bruto["intangible"][m] - amort_acum["intangible"][m]
            b["material"] = coste_bruto["material"][m] - amort_acum["material"][m]
        b["fianza"] = S["fianza"]
        b["clientes"] = (P["fact_base"][m] + P["fact_iva"][m]) if m >= 1 else 0.0
        b["hp_deudora_iva"] = -iva_saldo[m] if iva_saldo[m] < 0 else 0.0
        b["tesoreria"] = caja[m]
        b["total_activo"] = (b["intangible"] + b["material"] + b["fianza"] + b["clientes"]
                             + b["hp_deudora_iva"] + b["tesoreria"])
        ronda_hecha = bool(R["mes"]) and R["mes"] <= m
        b["capital"] = S["capital_fundadores"] + (nuevas if ronda_hecha else 0)
        b["prima"] = S["prima_fundadores"] + ((R["importe"] - nuevas) if ronda_hecha else 0.0)
        b["reservas"] = -S["gastos_constitucion"] - (R["gastos"] if ronda_hecha else 0.0)
        a = anio_de(m) if m >= 1 else 0
        b["anteriores"] = sum(resultados_anuales.get(x, 0.0) for x in range(1, a)) if m >= 1 else 0.0
        if m >= 1:
            meses_a = range((a - 1) * 12 + 1, m + 1)
            b["resultado"] = sum(P["ing_impl"][x] + P["ing_lic"][x] - P["sueldos"][x] - P["cargas"][x]
                                 - P["externos"][x] - P["comision"][x] - P["imprevistos"][x] - P["amort"][x]
                                 - P["financieros"][x] - P["is"][x] for x in meses_a)
        else:
            b["resultado"] = 0.0
        b["total_pn"] = b["capital"] + b["prima"] + b["reservas"] + b["anteriores"] + b["resultado"]
        # patrimonio neto a efectos del art. 363 LSC: con el préstamo participativo
        # (art. 20.1.d del Real Decreto-ley 7/1996)
        b["pn_mercantil"] = b["total_pn"] + enisa_pend[m]
        b["umbral_363"] = b["capital"] / 2.0
        e_corto, b_corto = corto_plazo(m)
        b["enisa_lp"] = enisa_pend[m] - e_corto
        b["banco_lp"] = banco_pend[m] - b_corto
        b["total_pnc"] = b["enisa_lp"] + b["banco_lp"]
        b["deuda_cp"] = e_corto + b_corto + intereses_pend[m]
        anticipos = diferidos = 0.0
        for c in clientes:
            firma = c["gm"] - 1
            if 1 <= firma <= m < c["gm"]:
                anticipos += S["precio_implantacion"] * S["parte_implantacion_firma"]
            fact = 0.0
            mm = c["gm"]
            while mm <= m:
                if c["fin"] is not None and mm >= c["fin"]:
                    break
                fact += precio_licencia_anio(S, (mm - c["gm"]) // 12)
                mm += 12
            rec = 0.0
            for x in range(c["gm"], m + 1):
                if c["fin"] is not None and x >= c["fin"]:
                    break
                rec += precio_licencia_anio(S, (x - c["gm"]) // 12) / 12.0
            diferidos += fact - rec
        b["anticipos"] = anticipos
        b["periodificacion"] = diferidos
        b["hp_acreedora"] = (iva_saldo[m] if iva_saldo[m] > 0 else 0.0) + irpf_pend[m] + is_pend[m]
        b["ss_acreedora"] = ss_pend[m]
        b["hp_ss"] = b["hp_acreedora"] + b["ss_acreedora"]
        b["total_pc"] = b["deuda_cp"] + b["anticipos"] + b["periodificacion"] + b["hp_ss"]
        b["total_pn_pasivo"] = b["total_pn"] + b["total_pnc"] + b["total_pc"]
        return b

    balances = {m: balance(m) for m in range(0, N + 1)}

    # ENISA presta como mucho los fondos propios; la solicitud se estudia con las cuentas
    # de 2026, es decir, con el balance del momento 0
    assert balances[0]["total_pn"] >= E["importe"], "ENISA: fondos propios por debajo del préstamo"
    comprobar_cuadre(balances)

    res = {
        "S": S, "N": N, "P": P, "caja": caja, "balances": balances, "clientes": clientes,
        "detalle_ext": detalle_ext, "personas_mes": personas_mes,
        "resultados_anuales": resultados_anuales, "bai_anual": bai_anual, "is_anual": is_anual,
        "bin_final": bin_disponible, "nuevas_participaciones": nuevas,
        "enisa_pend": enisa_pend, "banco_pend": banco_pend,
    }
    comprobar_encadenado(res)
    return res


def comprobar_cuadre(balances):
    """Prueba de cuadre: en cada cierre, el activo es igual al patrimonio neto más el pasivo, sumando
    otra vez las partidas una a una. Falla si no cuadra."""
    for m, b in balances.items():
        activo = b["intangible"] + b["material"] + b["fianza"] + b["clientes"] + b["hp_deudora_iva"] + b["tesoreria"]
        pn = b["capital"] + b["prima"] + b["reservas"] + b["anteriores"] + b["resultado"]
        pasivo = (b["enisa_lp"] + b["banco_lp"] + b["deuda_cp"] + b["anticipos"] + b["periodificacion"]
                  + b["hp_acreedora"] + b["ss_acreedora"])
        dif = activo - pn - pasivo
        assert abs(dif) < 0.01, f"El balance no cuadra en el mes {m}: diferencia {dif:.4f}"


def comprobar_encadenado(res):
    """Tesorería y resultados encadenados entre años; falla si algo no cuadra."""
    B = res["balances"]
    anios = res["N"] // 12
    for a in range(1, anios + 1):
        t = tesoreria_anual(res, a)
        y = pyg_anual(res, a)
        b_fin = B[a * 12]
        assert abs(b_fin["resultado"] - y["resultado"]) < 0.01, f"resultado del año {a}"
        assert abs(b_fin["tesoreria"] - t["final"]) < 0.01, f"tesorería del año {a}"
        if a > 1:
            b_ant = B[(a - 1) * 12]
            assert abs(b_fin["anteriores"] - (b_ant["anteriores"] + b_ant["resultado"])) < 0.01, \
                f"resultados anteriores del año {a}"
            assert abs(t["inicial"] - tesoreria_anual(res, a - 1)["final"]) < 0.01, f"tesorería inicial del año {a}"
    assert abs(B[0]["tesoreria"] - tesoreria_anual(res, 1)["inicial"]) < 0.01


# ---------------------------------------------------------------------------
# Agregados
# ---------------------------------------------------------------------------
def suma(lista, a):
    return sum(lista[m] for m in range((a - 1) * 12 + 1, a * 12 + 1))


def pyg_anual(res, a):
    P = res["P"]
    d = {"impl": suma(P["ing_impl"], a), "lic": suma(P["ing_lic"], a)}
    d["ingresos"] = d["impl"] + d["lic"]
    d["sueldos"] = suma(P["sueldos"], a)
    d["cargas"] = suma(P["cargas"], a)
    d["ext"] = {k: suma(v, a) for k, v in res["detalle_ext"].items()}
    d["comision"] = suma(P["comision"], a)
    d["imprevistos"] = suma(P["imprevistos"], a)
    d["gastos_expl"] = d["sueldos"] + d["cargas"] + sum(d["ext"].values()) + d["comision"] + d["imprevistos"]
    d["ebitda"] = d["ingresos"] - d["gastos_expl"]
    d["amort"] = suma(P["amort"], a)
    d["ebit"] = d["ebitda"] - d["amort"]
    d["financieros"] = suma(P["financieros"], a)
    d["bai"] = d["ebit"] - d["financieros"]
    d["is"] = suma(P["is"], a)
    d["resultado"] = d["bai"] - d["is"]
    return d


def tesoreria_anual(res, a):
    P = res["P"]
    t = {"inicial": res["caja"][(a - 1) * 12], "cobros": suma(P["cobros"], a)}
    t["personal"] = suma(P["pago_nomina"], a) + suma(P["pago_ss"], a) + suma(P["pago_irpf"], a)
    t["proveedores"] = suma(P["pago_externos"], a)
    t["iva_is"] = suma(P["pago_iva"], a) + suma(P["pago_is"], a)
    t["explotacion"] = t["cobros"] - t["personal"] - t["proveedores"] - t["iva_is"]
    t["inversion"] = -suma(P["pago_inversion"], a)
    t["capital"] = suma(P["entrada_capital"], a) - suma(P["pago_gastos_ronda"], a)
    t["prestamos"] = suma(P["entrada_prestamo"], a)
    t["intereses"] = -suma(P["pago_intereses"], a)
    t["principal"] = -suma(P["pago_principal"], a)
    t["financiacion"] = t["capital"] + t["prestamos"] + t["intereses"] + t["principal"]
    t["final"] = res["caja"][a * 12]
    assert abs(t["inicial"] + t["explotacion"] + t["inversion"] + t["financiacion"] - t["final"]) < 0.01
    return t


def activos_al_final(res, m):
    return sum(1 for c in res["clientes"] if c["gm"] <= m and (c["fin"] is None or c["fin"] > m))


def caja_minima(res, desde=1, hasta=None):
    hasta = hasta or res["N"]
    m = min(range(desde, hasta + 1), key=lambda x: res["caja"][x])
    return res["caja"][m], m


def holgura_363(res, hasta=None):
    """Mes con menos margen sobre la mitad del capital y ese margen (puede ser negativo)."""
    hasta = hasta if hasta is not None else res["N"]
    B = res["balances"]
    m = min(range(0, hasta + 1), key=lambda x: B[x]["pn_mercantil"] - B[x]["umbral_363"])
    return B[m]["pn_mercantil"] - B[m]["umbral_363"], m


def pn_minimo(res, hasta=None):
    hasta = hasta if hasta is not None else res["N"]
    B = res["balances"]
    m = min(range(0, hasta + 1), key=lambda x: B[x]["pn_mercantil"])
    return B[m]["pn_mercantil"], m


def punto_muerto(res, a=3):
    """Licencias cuyo margen paga la estructura del año a (como clientes_equilibrio de la v10)."""
    S = res["S"]
    y = pyg_anual(res, a)
    fijos = y["sueldos"] + y["cargas"] + sum(y["ext"].values())
    k, imp = S["comision_socio"], S["imprevistos"]
    lic_eq = (fijos * (1 + imp) + y["amort"] + y["financieros"]) / (1 - k * (1 + imp))
    return lic_eq / S["precio_licencia"], lic_eq


def van(res, tasa=TASA_VAN):
    """VAN de los flujos libres (explotación más inversiones) de todos los años, sin valor residual."""
    anios = res["N"] // 12
    flujos = []
    for a in range(1, anios + 1):
        t = tesoreria_anual(res, a)
        flujos.append(t["explotacion"] + t["inversion"])
    return sum(f / (1 + tasa) ** (i + 1) for i, f in enumerate(flujos)), flujos


def sueldo_renunciado(S):
    """Sueldo bruto de los fundadores al que renuncian, en euros (año 1)."""
    bruto_fund = sum(b for (_, b, desde) in S["plantilla"] if desde == 1)
    parte = [0.0] * (S["meses"] + 2)
    for (d, h, p) in S.get("recortes_fundadores", []):
        for m in range(d, min(h, S["meses"]) + 1):
            parte[m] = max(parte[m], p)
    return sum(bruto_fund / 12.0 * p for p in parte)


def cumple(r):
    """Caja y patrimonio del art. 363 en positivo todos los meses del horizonte."""
    h, mh = holgura_363(r)
    c, mc = caja_minima(r)
    return h >= -1e-6 and c >= -1e-6


def rebaja_minima(S, hasta_mes):
    """Parte mínima del sueldo de los fundadores a la que hay que renunciar en los meses 1 a
    hasta_mes para que caja y patrimonio cumplan en los 36 y en los 60 meses."""
    def prueba(p):
        s = copy.deepcopy(S)
        s["recortes_fundadores"] = [(1, hasta_mes, p)]
        return cumple(simular(s)) and cumple(simular(a_60_meses(s)))
    if not prueba(1.0):
        return None
    lo, hi = 0.0, 1.0
    if prueba(0.0):
        return 0.0
    for _ in range(40):
        mid = (lo + hi) / 2
        if prueba(mid):
            hi = mid
        else:
            lo = mid
    return hi


def enisa_maximo(S, paso=500.0):
    """Importe redondo más alto de ENISA (en múltiplos de paso) que no pasa de los fondos propios
    del momento 0 y con el que caja y patrimonio siguen cumpliendo."""
    r = simular(S)
    tope = math.floor(r["balances"][0]["total_pn"] / paso) * paso
    imp = tope
    while imp >= 25000.0:
        s = copy.deepcopy(S)
        s["enisa"]["importe"] = imp
        try:
            if cumple(simular(s)) and cumple(simular(a_60_meses(s))):
                return imp
        except AssertionError:
            pass
        imp -= paso
    return None


# ---------------------------------------------------------------------------
# Escenarios de sensibilidad
# ---------------------------------------------------------------------------
def escenario_retraso(S, meses=6, recorte=None):
    S2 = copy.deepcopy(S)
    S2["puestas_en_marcha"] = [g + meses for g in S["puestas_en_marcha"]]
    S2["bajas"] = [(g + meses, f + meses) for (g, f) in S["bajas"]]
    if recorte:
        S2["recortes_fundadores"] = list(S.get("recortes_fundadores", [])) + [recorte]
    return S2


def escenario_menos_clientes(S, menos=2):
    S2 = copy.deepcopy(S)
    a1 = [g for g in S["puestas_en_marcha"] if g <= 12]
    a2 = [g for g in S["puestas_en_marcha"] if 12 < g <= 24]
    a3 = [g for g in S["puestas_en_marcha"] if g > 24]
    S2["puestas_en_marcha"] = a1 + a2[:len(a2) - menos] + a3[:len(a3) - menos]
    S2["bajas"] = [(g, f) for (g, f) in S["bajas"] if g in S2["puestas_en_marcha"]]
    return S2


def escenario_sin_ronda(S):
    S2 = copy.deepcopy(S)
    S2["ronda"] = {"mes": None, "importe": 0.0, "participacion": 0.25, "gastos": 0.0}
    return S2


def escenario_enisa(S, mes):
    """ENISA llega en otro mes; el medio sueldo de los fundadores se queda en los meses 1 a 5."""
    S2 = copy.deepcopy(S)
    S2["enisa"]["mes_cobro"] = mes
    return S2


def escenario_enisa_importe(S, importe):
    S2 = copy.deepcopy(S)
    S2["enisa"]["importe"] = importe
    return S2


def escenario_enisa_denegado(S):
    S2 = copy.deepcopy(S)
    S2["enisa"]["importe"] = 0.0
    return S2


def escenario_red(S, ct, senales):
    S2 = copy.deepcopy(S)
    S2["precio_licencia"] = licencia(ct, senales)
    return S2


def escenario_mas_clientes(S):
    """Optimista: dos clientes nuevos más en los años 2 y 3, en los meses que quita el pesimista
    (22 y 24; 35 y 36), con el mismo ingeniero (decisión 12)."""
    S2 = copy.deepcopy(S)
    S2["puestas_en_marcha"] = sorted(S["puestas_en_marcha"] + [22, 24, 35, 36])
    return S2


def nota_contrataciones(S):
    """Texto con los meses de contratación que cambian frente a los previstos (13 y 25)."""
    meses = meses_contratacion(S)
    nombres = {13: "el ingeniero", 25: "el comercial"}
    partes = [(nombres.get(m, f"la contratación del mes {m}"), real) for m, real in sorted(meses.items()) if real != m]
    if not partes:
        return ""
    txt = []
    for i, (n, real) in enumerate(partes):
        if not real:
            txt.append(f"{n} no entra")
        else:
            txt.append(f"{n} entra en el mes {real}" if i == 0 else f"{n} en el {real}")
    return " (" + " y ".join(txt) + ")"


def lista_escenarios(S):
    if S["nombre"] == "v11":
        retraso6, retraso12 = escenario_retraso(S, 6), escenario_retraso(S, 12)
        menos = escenario_menos_clientes(S, 2)
        return [
            ("Moderado: el plan", S),
            ("Pesimista: dos clientes nuevos menos en los años 2 y 3" + nota_contrataciones(menos), menos),
            ("Optimista: dos clientes nuevos más en los años 2 y 3, con el mismo ingeniero", escenario_mas_clientes(S)),
            ("Los contratos llegan seis meses más tarde" + nota_contrataciones(retraso6), retraso6),
            ("Lo mismo, con medio sueldo de los fundadores hasta el mes 11", escenario_retraso(S, 6, recorte=(1, 11, 0.5))),
            ("Los contratos llegan un año más tarde" + nota_contrataciones(retraso12), retraso12),
            ("Sin la ronda", escenario_sin_ronda(S)),
            ("ENISA llega dos meses más tarde (mes 8)", escenario_enisa(S, MES_ENISA_V11 + 2)),
            ("ENISA presta solo el mínimo de la línea (25.000 €)", escenario_enisa_importe(S, 25000.0)),
            ("ENISA deniega el préstamo", escenario_enisa_denegado(S)),
            ("Clientes más pequeños: redes de 100 CT y 6.000 señales", escenario_red(S, 100, 6000)),
        ]
    lista = [
        ("Base: el plan", S),
        ("Los contratos llegan seis meses más tarde", escenario_retraso(S, 6)),
    ]
    if S["nombre"] == "v11":
        lista.append(("Lo mismo, con medio sueldo de los fundadores hasta el mes 11",
                      escenario_retraso(S, 6, recorte=(1, 11, 0.5))))
    else:
        lista.append(("Lo mismo, con los sueldos de los fundadores a la mitad de los meses 4 a 11",
                      escenario_retraso(S, 6, recorte=(4, 11, 0.5))))
    lista += [
        ("Los contratos llegan un año más tarde", escenario_retraso(S, 12)),
        ("Dos clientes nuevos menos en los años 2 y 3", escenario_menos_clientes(S, 2)),
        ("Sin la ronda", escenario_sin_ronda(S)),
    ]
    if S["nombre"] == "v11":
        lista += [
            ("ENISA llega dos meses más tarde (mes 8)", escenario_enisa(S, MES_ENISA_V11 + 2)),
            ("ENISA presta solo el mínimo de la línea (25.000 €)", escenario_enisa_importe(S, 25000.0)),
            ("ENISA deniega el préstamo", escenario_enisa_denegado(S)),
            ("Clientes más pequeños: redes de 100 CT y 6.000 señales", escenario_red(S, 100, 6000)),
        ]
    return lista


def escenarios(S):
    # en los escenarios el art. 363 y la caja se miden, no se exigen
    return [(nombre, simular(s)) for nombre, s in lista_escenarios(S)]


# ---------------------------------------------------------------------------
# Tablas en el formato de la memoria
# ---------------------------------------------------------------------------
def tabla(titulo, cabecera, filas, anchos=None):
    out = [f"T {titulo}", "H " + " | ".join(cabecera)]
    if anchos:
        out.append("W " + " | ".join(str(w) for w in anchos))
    for f in filas:
        assert len(f) == len(cabecera), (titulo, f)
        for celda in f:
            assert "|" not in str(celda)
        out.append("R " + " | ".join(str(c) for c in f))
    return "\n".join(out)


PALABRAS = {1: "uno", 2: "dos", 3: "tres", 4: "cuatro", 5: "cinco", 6: "seis", 7: "siete", 8: "ocho", 9: "nueve"}


def construir_tablas(res, escen):
    S = res["S"]
    tablas = {}
    y = {a: pyg_anual(res, a) for a in (1, 2, 3)}
    t = {a: tesoreria_anual(res, a) for a in (1, 2, 3)}
    B = res["balances"]
    v11 = S["nombre"] == "v11"

    # --- presupuesto del primer año
    e1 = y[1]["ext"]
    fund = [p for p in S["plantilla"] if p[2] == 1]
    meses_lic_a1 = sum(13 - g for g in S["puestas_en_marcha"] if g <= 12)
    pilotos_a1 = sum(1 for g in S["puestas_en_marcha"] if g <= 12)
    assert abs(y[1]["lic"] - meses_lic_a1 * S["precio_licencia"] / 12) < 0.01
    personal_a1 = y[1]["sueldos"] + y[1]["cargas"]
    rec = S.get("recortes_fundadores", [])
    nota_sueldos = f"Los dos fundadores, con {fmt(fund[0][1])} y {fmt(fund[1][1])} € brutos al año"
    if rec:
        d, h, p = rec[0]
        nota_sueldos += (f", a medio sueldo los meses {d} a {h}, hasta que llega ENISA" if p == 0.5
                         else f", con un {fmt_pct(p, 0)} menos los meses {d} a {h}")
    if v11:
        nota_fin = (f"Intereses del banco ({fmt_pct(S['banco']['interes'], 1)}) todo el año y de ENISA "
                    f"({fmt_pct(S['enisa']['interes'], 0)}) desde el mes {S['enisa']['mes_cobro']}, y la comisión de "
                    "apertura de ENISA")
    else:
        nota_fin = (f"Intereses de ENISA ({fmt_pct(S['enisa']['interes'], 0)}) y del banco "
                    f"({fmt_pct(S['banco']['interes'], 1)}) y comisión de apertura de ENISA")
    filas = [
        ["Implantaciones", fmt(y[1]["impl"]), f"{PALABRAS[pilotos_a1].capitalize()} pilotos a "
         f"{fmt(S['precio_implantacion'])} €"],
        ["Licencias", fmt(y[1]["lic"]), f"{PALABRAS[pilotos_a1].capitalize()} licencias de {fmt(S['precio_licencia'])} € "
         f"al año, con {PALABRAS[meses_lic_a1 // pilotos_a1]} meses de servicio de media en 2027"],
        ["Total ingresos", fmt(y[1]["ingresos"]), "Implantaciones y licencias del año"],
        ["Sueldos y salarios", fmt(y[1]["sueldos"]), nota_sueldos],
        ["Cargas sociales", fmt(y[1]["cargas"]), f"{fmt_pct(S['cargas_sociales'], 0)} de los sueldos"],
        ["Nube", fmt(e1["Nube (cómputo, almacenamiento, demos)"]), "Cómputo, almacenamiento, demos y el plan "
         f"comercial de la previsión, {fmt(e1['Nube (cómputo, almacenamiento, demos)'] / 12)} € al mes"],
        ["Herramientas de desarrollo", fmt(e1["Herramientas de desarrollo"]),
         f"{fmt(S['herramientas_persona_mes'])} € por persona y mes"],
        ["Asesoría jurídica", fmt(e1["Asesoría jurídica"]), "Contratos con el socio y con los clientes, encargo de "
         "tratamiento y análisis del AI Act"],
        ["Gestoría", fmt(e1["Gestoría"]), f"{fmt(e1['Gestoría'] / 12)} € al mes"],
        ["Actividad comercial con el socio", fmt(e1["Actividad comercial y ferias"]),
         "Viajes y material de las demos con el socio; lo que le paguemos el año 1 sale de aquí"],
        ["Equipos menores, viajes y otros", fmt(e1["Equipos menores, viajes y otros"]),
         "Material de oficina y desplazamientos"],
        ["Coworking", fmt(e1["Coworking"]), f"{fmt(S['coworking_persona_mes'])} € por persona y mes"],
        ["Seguro de responsabilidad civil", fmt(e1["Seguros (responsabilidad civil)"]),
         "El que pedirá cualquier distribuidora en el contrato"],
        ["Imprevistos", fmt(y[1]["imprevistos"]), f"{fmt_pct(S['imprevistos'], 0)} de los gastos anteriores"],
        ["Total gastos de explotación", fmt(y[1]["gastos_expl"]),
         f"El {fmt_pct(personal_a1 / y[1]['gastos_expl'], 0)} es personal"],
        ["Resultado bruto de explotación", fmt(y[1]["ebitda"]), "Ingresos menos gastos de explotación"],
        ["Amortizaciones", fmt(-y[1]["amort"]), "Equipos y aplicaciones en cuatro años; marca y mobiliario en diez"],
        ["Gastos financieros", fmt(-y[1]["financieros"]), nota_fin],
        ["Impuesto sobre Sociedades", fmt(-y[1]["is"]), "Año con pérdidas"],
        ["Resultado neto del año 1", fmt(y[1]["resultado"]),
         "Se compensa en el Impuesto sobre Sociedades con los beneficios de los años siguientes"],
    ]
    tablas["presupuesto_a1"] = tabla("Presupuesto del primer año (2027), en euros",
                                     ["Concepto", "Importe", "Cómo lo calculamos"], filas, [4, 2, 6])

    # --- cuenta de explotación
    def fila(nombre, clave=None, f=None, signo=1):
        return [nombre] + [fmt(signo * (f(y[a]) if f else y[a][clave])) for a in (1, 2, 3)]

    ext_orden = ["Nube (cómputo, almacenamiento, demos)", "Herramientas de desarrollo", "Asesoría jurídica",
                 "Gestoría", "Actividad comercial y ferias", "Equipos menores, viajes y otros", "Coworking",
                 "Seguros (responsabilidad civil)"]
    nombres_ext = {"Nube (cómputo, almacenamiento, demos)": "Nube (cómputo, almacenamiento y demos)",
                   "Seguros (responsabilidad civil)": "Seguro de responsabilidad civil"}
    nuevos = [sum(1 for g in S["puestas_en_marcha"] if (a - 1) * 12 < g <= a * 12) for a in (1, 2, 3)]
    filas = [
        ["Clientes nuevos en el año"] + [str(n) for n in nuevos],
        ["Clientes activos al final del año"] + [str(activos_al_final(res, a * 12)) for a in (1, 2, 3)],
        ["Licencias anuales en cartera al final del año"]
        + [fmt(activos_al_final(res, a * 12) * S["precio_licencia"]) for a in (1, 2, 3)],
        fila("Ingresos por implantación", "impl"),
        fila("Ingresos por licencias", "lic"),
        fila("Total ingresos", "ingresos"),
        fila("Sueldos y salarios", "sueldos", signo=-1),
        fila("Cargas sociales", "cargas", signo=-1),
    ]
    for k in ext_orden:
        filas.append(fila(nombres_ext.get(k, k), f=lambda d, k=k: d["ext"][k], signo=-1))
    filas += [
        fila("Comisión del socio (10 % de las licencias)", "comision", signo=-1),
        fila("Imprevistos (5 %)", "imprevistos", signo=-1),
        fila("Total gastos de explotación", "gastos_expl", signo=-1),
        fila("Resultado bruto de explotación", "ebitda"),
        fila("Amortizaciones", "amort", signo=-1),
        fila("Resultado de explotación", "ebit"),
        fila("Gastos financieros", "financieros", signo=-1),
        fila("Resultado antes de impuestos", "bai"),
        fila("Impuesto sobre Sociedades", "is", signo=-1),
        fila("Resultado del ejercicio", "resultado"),
        fila("Cash flow (resultado más amortizaciones)", f=lambda d: d["resultado"] + d["amort"]),
    ]
    tablas["pyg"] = tabla("Cuenta de explotación a tres años, en euros",
                          ["Concepto", "Año 1", "Año 2", "Año 3"], filas, [6, 2, 2, 2])

    # --- balances
    def bfila(nombre, clave, momentos=(0, 12, 24, 36)):
        return [nombre] + [fmt(B[m][clave]) for m in momentos]
    filas = [
        bfila("Inmovilizado intangible (marca y aplicaciones)", "intangible"),
        bfila("Inmovilizado material (equipos y mobiliario)", "material"),
        ["Total activo no corriente"] + [fmt(B[m]["intangible"] + B[m]["material"]) for m in (0, 12, 24, 36)],
        bfila("Fianza del coworking", "fianza"),
        bfila("Clientes", "clientes"),
        bfila("Hacienda Pública deudora por IVA", "hp_deudora_iva"),
        bfila("Tesorería", "tesoreria"),
        ["Total activo corriente"] + [fmt(B[m]["fianza"] + B[m]["clientes"] + B[m]["hp_deudora_iva"]
                                          + B[m]["tesoreria"]) for m in (0, 12, 24, 36)],
        bfila("Total activo", "total_activo"),
        bfila("Capital social", "capital"),
        bfila("Prima de asunción", "prima"),
        bfila("Reservas (gastos de constitución y de la ronda)", "reservas"),
        bfila("Resultados de ejercicios anteriores", "anteriores"),
        bfila("Resultado del ejercicio", "resultado"),
        bfila("Total patrimonio neto", "total_pn"),
        bfila("Préstamo de ENISA a largo plazo", "enisa_lp"),
        bfila("Préstamo bancario a largo plazo", "banco_lp"),
        bfila("Total pasivo no corriente", "total_pnc"),
        bfila("Cuotas de los préstamos que vencen en el año siguiente", "deuda_cp"),
        ["Ingresos anticipados (licencias y anticipos facturados)"]
        + [fmt(B[m]["periodificacion"] + B[m]["anticipos"]) for m in (0, 12, 24, 36)],
        bfila("Hacienda Pública y Seguridad Social acreedoras", "hp_ss"),
        bfila("Total pasivo corriente", "total_pc"),
        bfila("Total patrimonio neto y pasivo", "total_pn_pasivo"),
    ]
    tablas["balances"] = tabla("Balances de situación de los momentos 0 a 3, en euros",
                               ["Partida", "Momento 0", "Momento 1", "Momento 2", "Momento 3"],
                               filas, [6, 2, 2, 2, 2])

    # --- tesorería del año 1, mes a mes
    P = res["P"]
    filas = []
    for m in range(1, 13):
        cobros = P["cobros"][m]
        personal = P["pago_nomina"][m] + P["pago_ss"][m] + P["pago_irpf"][m]
        prov = P["pago_externos"][m] + P["pago_inversion"][m]
        iva_is = P["pago_iva"][m] + P["pago_is"][m]
        finan = (P["entrada_capital"][m] + P["entrada_prestamo"][m] - P["pago_gastos_ronda"][m]
                 - P["pago_intereses"][m] - P["pago_principal"][m])
        neto = cobros - personal - prov - iva_is + finan
        assert abs(res["caja"][m - 1] + neto - res["caja"][m]) < 0.01
        filas.append([str(m), fmt(cobros), fmt(-personal), fmt(-prov), fmt(-iva_is), fmt(finan), fmt(neto),
                      fmt(res["caja"][m]), fmt(B[m]["pn_mercantil"])])
    tablas["tesoreria_mensual"] = tabla(
        "Tesorería del año 1 mes a mes, en euros (cobros y pagos con IVA), y patrimonio neto con el préstamo "
        "participativo (art. 363 LSC)",
        ["Mes", "Cobros", "Personal y Seguridad Social", "Proveedores", "IVA", "Financiación",
         "Saldo del mes", "Tesorería final", "Patrimonio neto (art. 363)"], filas, [1, 2, 2, 2, 2, 2, 2, 2, 2])

    # --- tesorería de los tres años
    def tfila(nombre, clave, signo=1):
        return [nombre] + [fmt(signo * t[a][clave]) for a in (1, 2, 3)]
    filas = [
        tfila("Tesorería inicial", "inicial"),
        tfila("Cobros de clientes", "cobros"),
        tfila("Pagos de personal y Seguridad Social", "personal", -1),
        tfila("Pagos a proveedores", "proveedores", -1),
        tfila("IVA e Impuesto sobre Sociedades", "iva_is", -1),
        tfila("Flujo de caja de la explotación", "explotacion"),
        tfila("Inversiones", "inversion"),
        tfila("Ronda de capital, neta de gastos", "capital"),
    ]
    if any(t[a]["prestamos"] for a in (1, 2, 3)):
        filas.append(tfila("Préstamo de ENISA cobrado", "prestamos"))
    filas += [
        tfila("Intereses pagados", "intereses"),
        tfila("Devolución de préstamos", "principal"),
        tfila("Tesorería final", "final"),
    ]
    tablas["tesoreria_anual"] = tabla("Tesorería de los tres años, en euros",
                                      ["Concepto", "Año 1", "Año 2", "Año 3"], filas, [6, 2, 2, 2])

    # --- estructura de financiación
    filas = []
    mc = S["enisa"]["mes_cobro"]
    momentos = [("Momento 0", 0)]
    if mc > 0:
        momentos.append((f"Mes {mc - 1}, antes de ENISA", mc - 1))
    momentos += [("Mes 11, antes de la ronda", 11), ("Momento 1", 12), ("Momento 2", 24), ("Momento 3", 36)]
    for etiqueta, m in momentos:
        pr = B[m]["total_pn"]
        aj = res["enisa_pend"][m] + res["banco_pend"][m]
        reparto = ("Patrimonio neto negativo" if pr < 0
                   else fmt_pct(pr / (pr + aj), 0) + " / " + fmt_pct(aj / (pr + aj), 0))
        filas.append([etiqueta, fmt(pr), fmt(aj), reparto, fmt(B[m]["pn_mercantil"])])
    tablas["estructura"] = tabla("Fondos propios frente a deuda financiera, en euros",
                                 ["Momento", "Patrimonio neto", "Préstamos (ENISA y banco)", "Propio / ajeno",
                                  "Patrimonio neto con el participativo (art. 363)"], filas, [3, 2, 2, 2, 3])

    # --- uso de la ronda
    a2, a3 = y[2], y[3]
    equipo_nuevo = 0.0
    for (_, bruto, desde) in S["plantilla"]:
        if desde > 12:
            meses_a2 = max(0, 24 - max(desde, 13) + 1)
            meses_a3 = max(0, 36 - max(desde, 25) + 1)
            equipo_nuevo += bruto / 12.0 * (meses_a2 + meses_a3) * (1 + S["cargas_sociales"])
    equipo_nuevo += sum(imp for (mes, _, imp, tipo) in S["inversiones"] if tipo == "equipo" and mes > 12)
    comercial = a2["ext"]["Actividad comercial y ferias"] + a3["ext"]["Actividad comercial y ferias"]
    legal = (a2["ext"]["Asesoría jurídica"] + a3["ext"]["Asesoría jurídica"] + S["ronda"]["gastos"]
             + sum(imp for (mes, _, imp, tipo) in S["inversiones"] if tipo == "marca"))
    imprev = a2["imprevistos"] + a3["imprevistos"]
    filas = [
        ["Equipo", "150.000", fmt(equipo_nuevo),
         "Sueldo, cargas y puesto del ingeniero desde el mes 13 y del responsable comercial desde el mes 25"],
        ["Comercial y ferias", "60.000", fmt(comercial),
         "Ferias y jornadas del sector, cuotas de asociaciones y viajes de demo de los años 2 y 3"],
        ["Legal y contingencia", "40.000", fmt(legal + imprev),
         f"Asesoría jurídica, marca en la EUIPO y gastos de la ronda ({fmt(legal)} €) e imprevistos del 5 % "
         f"({fmt(imprev)} €)"],
        ["Total", "250.000", fmt(equipo_nuevo + comercial + legal + imprev),
         "Lo que no cubren los ingresos queda en caja como colchón"],
    ]
    tablas["uso_ronda"] = tabla("Uso de la ronda: destino previsto y lo que gasta el plan en los años 2 y 3, más los "
                                "gastos de la ronda, en euros",
                                ["Partida", "Destino previsto", "Gasto del plan", "Qué incluye"], filas, [2, 2, 2, 6])

    # --- sensibilidad
    filas = []
    for nombre, r in escen:
        cm, mm = caja_minima(r, 1, 36)
        h, mh = holgura_363(r, 36)
        filas.append([nombre, fmt(cm) + f" (mes {mm})", fmt(h + r["balances"][mh]["umbral_363"]) + f" (mes {mh})",
                      fmt(pyg_anual(r, 2)["resultado"]),
                      fmt(pyg_anual(r, 3)["ingresos"]) + f" ({activos_al_final(r, 36)})",
                      fmt(pyg_anual(r, 3)["resultado"]), fmt(r["caja"][36])])
    tablas["sensibilidad"] = tabla(
        "Sensibilidad del plan, en euros",
        ["Escenario", "Caja más baja (mes)", "Patrimonio neto más bajo con el participativo (mes)",
         "Resultado del año 2", "Ingresos del año 3 (clientes al final)", "Resultado del año 3",
         "Tesorería al final del año 3"], filas, [5, 2, 2, 2, 2, 2, 2])

    # --- tarifa
    filas = []
    redes = list(REDES_CALIBRACION.items()) + [("Red tipo del plan", RED_TIPO_PLAN)] + list(REDES_REFERENCIA.items())
    for nombre, red in redes:
        fijo, pct, psen = partes_licencia(red["ct"], red["senales"])
        lic = licencia(**red)
        sen = fmt(red["senales"])
        if nombre in REDES_CALIBRACION and SENALES_PROVISIONALES:
            sen = "unas " + sen + " (provisional)"
        nota = "Mínimo" if lic > fijo + pct + psen else ("Licencia de cada cliente del plan"
                                                          if nombre == "Red tipo del plan" else "Fijo + CT + señales")
        filas.append([nombre, fmt(red["ct"]), sen, fmt(pct), fmt(psen), fmt(lic), nota])
    tablas["tarifa"] = tabla(
        f"Licencia anual con la tarifa por tamaño de red ({texto_tarifa()}), en euros",
        ["Red", "CT", "Señales", "Parte de los CT", "Parte de las señales", "Licencia al año", "Qué se aplica"],
        filas, [3, 1, 2, 2, 2, 2, 3])

    # --- supuestos
    gm = S["puestas_en_marcha"]
    a1l = [g for g in gm if g <= 12]
    a2l = [g for g in gm if 12 < g <= 24]
    a3l = [g for g in gm if g > 24]
    pl = S["plantilla"]
    E, Bk, R = S["enisa"], S["banco"], S["ronda"]

    def lista_meses(v):
        v = [str(x) for x in v]
        return ", ".join(v[:-1]) + " y " + v[-1]
    if v11:
        precio = (f"{fmt(S['precio_implantacion'])} € de implantación. La licencia, por el tamaño de la red: "
                  f"{texto_tarifa()}. Cada cliente del plan es una red de {RED_TIPO_PLAN['ct']} CT y "
                  f"{fmt(RED_TIPO_PLAN['senales'])} señales, que paga {fmt(S['precio_licencia'])} € al año, sin subida "
                  "con el IPC.")
        precio_fuente = "M2 y Pitch v3 para la implantación; la tarifa, decisión nuestra calibrada con las tres redes del TFM (3.7.1)"
        plantilla = (f"Los dos fundadores ({fmt(pl[0][1])} y {fmt(pl[1][1])} € brutos al año), a medio sueldo los meses "
                     f"{RECORTE_FUNDADORES_V11[0]} a {RECORTE_FUNDADORES_V11[1]}, hasta que llega ENISA. Un ingeniero de "
                     f"implantación y soporte desde el mes {pl[2][2]} ({fmt(pl[2][1])} €) y un responsable comercial del "
                     f"sector desde el mes {pl[3][2]} ({fmt(pl[3][1])} €).")
        financiacion = (f"{fmt(S['capital_fundadores'] + S['prima_fundadores'])} € de los fundadores, la mitad cada uno "
                        f"({fmt(S['capital_fundadores'])} de capital y {fmt(S['prima_fundadores'])} de prima de asunción). "
                        f"Préstamo bancario de la línea del ICO de {fmt(Bk['importe'])} € a {Bk['plazo_meses'] // 12} años "
                        f"con {Bk['carencia_meses'] // 12} de carencia, en el momento 0. Préstamo participativo de ENISA de "
                        f"{fmt(E['importe'])} € a {E['plazo_meses'] // 12} años con {E['carencia_meses'] // 12} de carencia, "
                        f"cobrado en el mes {E['mes_cobro']}. Ronda de {fmt(R['importe'])} € por el 25 % en el mes "
                        f"{R['mes']}, con {fmt(R['gastos'])} € de gastos.")
        fin_fuente = ("Decisión nuestra; ENISA pide fondos propios al menos iguales al préstamo y las cuentas del último "
                      "ejercicio depositadas; la ronda, del Pitch v3")
    else:
        precio = (f"{fmt(S['precio_implantacion'])} € de implantación y {fmt(S['precio_licencia'])} € al año de "
                  "licencia, sin subida con el IPC.")
        precio_fuente = "M2 y Pitch v3"
        plantilla = (f"Los dos fundadores ({fmt(pl[0][1])} y {fmt(pl[1][1])} € brutos al año). Un ingeniero de "
                     f"implantación y soporte desde el mes {pl[2][2]} ({fmt(pl[2][1])} €) y un responsable comercial del "
                     f"sector desde el mes {pl[3][2]} ({fmt(pl[3][1])} €).")
        financiacion = (f"{fmt(S['capital_fundadores'] + S['prima_fundadores'])} € de los fundadores, la mitad cada uno "
                        f"({fmt(S['capital_fundadores'])} de capital y {fmt(S['prima_fundadores'])} de prima de asunción). "
                        f"Préstamo participativo de ENISA de {fmt(E['importe'])} € a {E['plazo_meses'] // 12} años con "
                        f"{E['carencia_meses'] // 12} de carencia. Préstamo bancario de la línea del ICO de "
                        f"{fmt(Bk['importe'])} € a {Bk['plazo_meses'] // 12} años con {Bk['carencia_meses'] // 12} de "
                        f"carencia. Ronda de {fmt(R['importe'])} € por el 25 % en el mes {R['mes']}, con "
                        f"{fmt(R['gastos'])} € de gastos.")
        fin_fuente = "Decisión nuestra; ENISA pide fondos propios al menos iguales al préstamo; la ronda, del Pitch v3"
    filas = [
        ["Calendario", "Constituimos la sociedad en diciembre de 2026 (momento 0); el año 1 es 2027, el 2 es 2028 "
         "y el 3 es 2029. El mes 1 es enero de 2027.", "Decisión nuestra"],
        ["Clientes nuevos", f"{len(a1l)} el año 1, {len(a2l)} el año 2 y {len(a3l)} el año 3.",
         "Pitch v3 para los años 1 y 2 (tres pilotos; de 6 a 8 clientes); el año 3, decisión nuestra para llegar "
         "al suelo de ingresos del Pitch"],
        ["Ritmo de las firmas", f"Puesta en marcha en los meses {lista_meses(a1l)} del año 1; en los meses "
         f"{lista_meses(a2l)} del año 2, y en los meses {lista_meses(a3l)} del año 3. Antes, un mes de demo gratuita; "
         "la firma, el mes anterior a la puesta en marcha.",
         f"Decisión nuestra; los pilotos suman {meses_lic_a1} meses de licencia el año 1, los mismos que el M2 "
         "(seis meses de media)"],
        ["Bajas", "Uno de los tres pilotos no renueva al cumplir el año (mes 21) y otro cliente se va al cumplir el "
         "suyo (mes 26): dos bajas entre los 17 clientes del plan.", "Decisión nuestra, prudente"],
        ["Precio", precio, precio_fuente],
        ["Facturación y cobro", "La implantación, mitad a la firma y mitad en la puesta en marcha; la licencia, por "
         "años adelantados. Todo se cobra a 30 días.", "Decisión nuestra"],
        ["Plantilla", plantilla, "M2 para los fundadores; las contrataciones, decisión nuestra"],
        ["Cargas sociales", f"{fmt_pct(S['cargas_sociales'], 0)} de los salarios brutos.", "M2"],
        ["Otros gastos", "El año 1, 24.000 € de nube, herramientas, asesoría, gestoría, actividad comercial y otros, "
         "más el coworking (250 € por persona y mes) y un seguro de responsabilidad civil; las herramientas, a 100 € por "
         "persona y mes. En los años 2 y 3 crecen con la plantilla y con la actividad comercial.",
         "M2; la fianza de dos meses del balance de momento 0; decisión nuestra"],
        ["Comisión del socio", f"{fmt_pct(S['comision_socio'], 0)} de las licencias desde el año 2, por instalar y dar "
         "el primer soporte. El año 1 va dentro de la partida de actividad comercial.",
         "Decisión nuestra, por negociar con SITEL"],
        ["Imprevistos", f"{fmt_pct(S['imprevistos'], 0)} de los gastos de explotación.", "Decisión nuestra"],
        ["Inversiones y amortización", "El activo del momento 0 (13.000 €), 2.500 € de puesto de trabajo por cada "
         "contratación y 1.000 € para la marca en la EUIPO. Amortización lineal: marca y mobiliario en diez años; "
         "aplicaciones y equipos informáticos en cuatro.",
         "Balance de momento 0 y su plantilla del año 1; las aplicaciones, decisión nuestra"],
        ["Financiación", financiacion, fin_fuente],
        ["Intereses", f"{fmt_pct(E['interes'], 0)} el de ENISA (tramo fijo y variable juntos) y "
         f"{fmt_pct(Bk['interes'], 1)} el bancario; comisión de apertura de ENISA del "
         f"{fmt_pct(E['comision_apertura'], 1)}.", "Decisión nuestra, prudente"],
        ["Impuestos", f"IVA del {fmt_pct(S['iva'], 0)} y retenciones de IRPF del {fmt_pct(S['retencion_irpf'], 0)}, que "
         "se liquidan cada trimestre; la Seguridad Social, el mes siguiente. Impuesto sobre Sociedades del "
         f"{fmt_pct(S['tipo_is'], 0)} (entidad de nueva creación), compensando las pérdidas anteriores y pagado en "
         "julio del año siguiente; el crédito fiscal no se lleva al balance.", "Ley 27/2014 y decisión nuestra"],
        ["Proveedores", "Se pagan al contado.", "Decisión nuestra, prudente"],
    ]
    tablas["supuestos"] = tabla("Supuestos del plan financiero", ["Concepto", "Supuesto", "De dónde sale"],
                                filas, [2, 6, 3])
    tablas.update(tablas_analisis(res))
    return tablas, y, t


# ---------------------------------------------------------------------------
# Análisis que cita el texto de 3.3, 3.8.3, 3.9 y 3.10.5
# ---------------------------------------------------------------------------
# Valor de un minuto de TIEPI de toda la red en el incentivo de calidad (memoria, 3.8.3) y
# tasa de retribución financiera de la CNMC para 2026 - 2031, con la que se descuenta el
# valor para el cliente.
MINUTO_TIEPI = {"Distribuidora 1": 12600.0, "Distribuidora 2": 14900.0}
TASA_CNMC = 0.0658
IPC_CONTRASTE = 0.03


def cascada(res, a):
    """Cuenta de resultados en cascada: coste de las ventas = nube, comisión del socio y sueldo y
    cargas del ingeniero de implantación (con su mes real de entrada)."""
    S = res["S"]
    y = pyg_anual(res, a)
    nube = y["ext"]["Nube (cómputo, almacenamiento, demos)"]
    ingeniero = 0.0
    for (puesto, bruto, desde) in S["plantilla"]:
        if puesto.startswith("Ingeniero"):
            meses = sum(1 for m in range((a - 1) * 12 + 1, a * 12 + 1) if m >= desde)
            ingeniero += bruto / 12.0 * meses * (1 + S["cargas_sociales"])
    cv = nube + y["comision"] + ingeniero
    return {"ventas": y["ingresos"], "nube": nube, "comision": y["comision"], "ingeniero": ingeniero,
            "coste_ventas": cv, "margen_bruto": y["ingresos"] - cv, "margen_bruto_pct": (y["ingresos"] - cv) / y["ingresos"],
            "gastos_operativos": y["gastos_expl"] - cv, "ebitda": y["ebitda"], "amort": y["amort"], "ebit": y["ebit"],
            "financieros": y["financieros"], "bai": y["bai"], "is": y["is"], "resultado": y["resultado"],
            "cash_flow": y["resultado"] + y["amort"]}


def punto_muerto_detalle(res, a):
    """Punto muerto con la fórmula del curso: costes fijos / margen de contribución de una licencia;
    ventas de equilibrio con la mezcla del año y margen de seguridad = (ventas - equilibrio) / ventas."""
    S = res["S"]
    y = pyg_anual(res, a)
    k, imp, L = S["comision_socio"], S["imprevistos"], S["precio_licencia"]
    fijos = (y["sueldos"] + y["cargas"] + sum(y["ext"].values())) * (1 + imp) + y["amort"] + y["financieros"]
    aporte = L * (1 - k * (1 + imp))
    variables = y["comision"] * (1 + imp)
    ventas_eq = fijos / ((y["ingresos"] - variables) / y["ingresos"])
    costes = y["gastos_expl"] + y["amort"] + y["financieros"]
    return {"costes_fijos": fijos, "aporte_licencia": aporte, "licencias": fijos / aporte, "euros": fijos / aporte * L,
            "ventas_equilibrio": ventas_eq, "ventas": y["ingresos"],
            "margen_seguridad": (y["ingresos"] - ventas_eq) / y["ingresos"], "peso_fijos": fijos / costes}


def fondo_maniobra(res, m):
    """Activo corriente sin la fianza menos pasivo corriente, y necesidades operativas de fondos."""
    b = res["balances"][m]
    ac = b["clientes"] + b["hp_deudora_iva"] + b["tesoreria"]
    nof = b["clientes"] + b["hp_deudora_iva"] - (b["anticipos"] + b["periodificacion"] + b["hp_ss"])
    return {"activo_corriente": ac, "pasivo_corriente": b["total_pc"], "fondo_maniobra": ac - b["total_pc"], "nof": nof}


def valor_cliente(lic, minuto, tasa=TASA_CNMC, implantacion=20000.0, anios=3):
    """Valor actual para el cliente de la implantación y tres años de licencia pagados por adelantado, y
    minutos de TIEPI al año que lo compensan (renta vencida de los mismos años)."""
    v = 1.0 / (1.0 + tasa)
    anticipada = sum(v ** t for t in range(anios))
    vencida = sum(v ** t for t in range(1, anios + 1))
    valor = implantacion + lic * anticipada
    return {"licencia": lic, "valor_actual": valor, "minutos": valor / vencida / minuto}


def inversor(S, clientes):
    veces = clientes * S["precio_licencia"] * MULTIPLO_VENTA * S["ronda"]["participacion"] / S["ronda"]["importe"]
    return {"clientes": clientes, "licencias": clientes * S["precio_licencia"],
            "valor_venta": clientes * S["precio_licencia"] * MULTIPLO_VENTA, "veces": veces,
            "tir": veces ** (1.0 / ANIOS_HASTA_VENTA) - 1}


def valor_ipc(S, tasa=IPC_CONTRASTE):
    """Ingresos por licencias de más en los años 2 y 3 si la licencia subiera cada año en su aniversario."""
    base = simular(S)
    s2 = copy.deepcopy(S)
    s2["ipc_licencia"] = tasa
    con = simular(s2)
    return {a: pyg_anual(con, a)["lic"] - pyg_anual(base, a)["lic"] for a in (1, 2, 3)}


def tablas_analisis(res):
    S = res["S"]
    tablas = {}
    c = {a: cascada(res, a) for a in (1, 2, 3)}
    filas = [["Ventas"] + [fmt(c[a]["ventas"]) for a in (1, 2, 3)],
             ["Nube"] + [fmt(-c[a]["nube"]) for a in (1, 2, 3)],
             ["Comisión del socio"] + [fmt(-c[a]["comision"]) for a in (1, 2, 3)],
             ["Ingeniero de implantación (sueldo y cargas)"] + [fmt(-c[a]["ingeniero"]) for a in (1, 2, 3)],
             ["Coste de las ventas"] + [fmt(-c[a]["coste_ventas"]) for a in (1, 2, 3)],
             ["Margen bruto (sobre las ventas)"] + [f"{fmt(c[a]['margen_bruto'])} ({fmt_pct(c[a]['margen_bruto_pct'], 0)})"
                                                    for a in (1, 2, 3)],
             ["Gastos operativos"] + [fmt(-c[a]["gastos_operativos"]) for a in (1, 2, 3)],
             ["Resultado bruto de explotación (EBITDA)"] + [fmt(c[a]["ebitda"]) for a in (1, 2, 3)],
             ["Amortizaciones"] + [fmt(-c[a]["amort"]) for a in (1, 2, 3)],
             ["Resultado de explotación (EBIT)"] + [fmt(c[a]["ebit"]) for a in (1, 2, 3)],
             ["Gastos financieros"] + [fmt(-c[a]["financieros"]) for a in (1, 2, 3)],
             ["Resultado antes de impuestos"] + [fmt(c[a]["bai"]) for a in (1, 2, 3)],
             ["Impuesto sobre Sociedades"] + [fmt(-c[a]["is"]) for a in (1, 2, 3)],
             ["Resultado del ejercicio"] + [fmt(c[a]["resultado"]) for a in (1, 2, 3)],
             ["Cash flow (resultado más amortizaciones)"] + [fmt(c[a]["cash_flow"]) for a in (1, 2, 3)]]
    tablas["cascada"] = tabla("Cuenta de explotación en cascada, en euros", ["Concepto", "Año 1", "Año 2", "Año 3"],
                              filas, [6, 2, 2, 2])

    p = {a: punto_muerto_detalle(res, a) for a in (2, 3)}
    filas = [["Costes fijos (estructura con imprevistos, amortizaciones e intereses)"] + [fmt(p[a]["costes_fijos"]) for a in (2, 3)],
             ["Margen de contribución de una licencia"] + [fmt(p[a]["aporte_licencia"]) for a in (2, 3)],
             ["Licencias de equilibrio"] + [fmt1(p[a]["licencias"]) for a in (2, 3)],
             ["En euros de licencias"] + [fmt(p[a]["euros"]) for a in (2, 3)],
             ["Ventas de equilibrio con la mezcla del año"] + [fmt(p[a]["ventas_equilibrio"]) for a in (2, 3)],
             ["Ventas del plan"] + [fmt(p[a]["ventas"]) for a in (2, 3)],
             ["Margen de seguridad"] + [fmt_pct(p[a]["margen_seguridad"], 1) for a in (2, 3)],
             ["Peso de los costes fijos en los costes del año"] + [fmt_pct(p[a]["peso_fijos"], 1) for a in (2, 3)]]
    tablas["punto_muerto"] = tabla("Punto muerto y margen de seguridad, en euros", ["Concepto", "Año 2", "Año 3"],
                                   filas, [6, 2, 2])

    f = {m: fondo_maniobra(res, m) for m in (0, 12, 24, 36)}
    peor = min(range(1, 12), key=lambda m: fondo_maniobra(res, m)["fondo_maniobra"])
    filas = [["Activo corriente (sin la fianza)"] + [fmt(f[m]["activo_corriente"]) for m in (0, 12, 24, 36)],
             ["Pasivo corriente"] + [fmt(f[m]["pasivo_corriente"]) for m in (0, 12, 24, 36)],
             ["Fondo de maniobra"] + [fmt(f[m]["fondo_maniobra"]) for m in (0, 12, 24, 36)],
             ["Necesidades operativas de fondos"] + [fmt(f[m]["nof"]) for m in (0, 12, 24, 36)]]
    tablas["fondo_maniobra"] = tabla(
        f"Fondo de maniobra y necesidades operativas de fondos, en euros (el peor mes antes de la ronda es el {peor}, "
        f"con {fmt(fondo_maniobra(res, peor)['fondo_maniobra'])})",
        ["Concepto", "Momento 0", "Momento 1", "Momento 2", "Momento 3"], filas, [6, 2, 2, 2, 2])

    filas = []
    for nombre, minuto in MINUTO_TIEPI.items():
        v = valor_cliente(licencia(**REDES_CALIBRACION[nombre]), minuto)
        filas.append([nombre, fmt(v["licencia"]), fmt(minuto), fmt(v["valor_actual"]), fmt(round(v["valor_actual"], -2)),
                      fmt1(v["minutos"])])
    tablas["valor_cliente"] = tabla(
        f"Valor actual para el cliente de la implantación y tres años de licencia por adelantado, al "
        f"{fmt_pct(TASA_CNMC, 2)}, en euros",
        ["Red", "Licencia al año", "Un minuto de TIEPI", "Valor actual", "Redondeado", "Minutos de TIEPI al año que lo compensan"],
        filas, [3, 2, 2, 2, 2, 2])

    filas = []
    for n in (20, 30, 45):
        i = inversor(S, n)
        filas.append([str(n), fmt(i["licencias"]), fmt(i["valor_venta"]), fmt1(i["veces"]), fmt_pct(i["tir"], 1)])
    tablas["inversor"] = tabla(
        f"Lo que recuperaría el inversor de la ronda si la empresa se vendiera a finales de 2031 por "
        f"{fmt1(MULTIPLO_VENTA)} veces sus licencias anuales, sin más rondas, en euros",
        ["Clientes", "Licencias anuales", "Precio de venta", "Veces lo invertido", "Rentabilidad al año"],
        filas, [2, 2, 2, 2, 2])

    if S["nombre"] == "v11":
        ipc = valor_ipc(S)
        tablas["ipc"] = tabla(
            f"Ingresos por licencias de más si la licencia subiera un {fmt_pct(IPC_CONTRASTE, 0)} en cada aniversario "
            "(el plan no lo cuenta), en euros",
            ["Concepto", "Año 1", "Año 2", "Año 3"], [["Licencias de más"] + [fmt(ipc[a]) for a in (1, 2, 3)]],
            [6, 2, 2, 2])
    return tablas


# ---------------------------------------------------------------------------
# Cifras clave, cifras del texto y comparación
# ---------------------------------------------------------------------------
def cifras_clave(r36, r60):
    c = {}
    for a in (1, 2, 3):
        y = pyg_anual(r36, a)
        c[f"ingresos_a{a}"] = y["ingresos"]
        c[f"resultado_a{a}"] = y["resultado"]
        c[f"bai_a{a}"] = y["bai"]
        c[f"caja_fin_a{a}"] = r36["caja"][a * 12]
        c[f"pn_a{a}"] = r36["balances"][a * 12]["total_pn"]
    for a in (4, 5):
        y = pyg_anual(r60, a)
        c[f"ingresos_a{a}"] = y["ingresos"]
        c[f"resultado_a{a}"] = y["resultado"]
        c[f"caja_fin_a{a}"] = r60["caja"][a * 12]
    pm, lic_eq = punto_muerto(r36, 3)
    c["punto_muerto_licencias"] = pm
    c["punto_muerto_euros"] = lic_eq
    c["licencias_cartera_a3"] = activos_al_final(r36, 36) * r36["S"]["precio_licencia"]
    cm, mm = caja_minima(r36)
    c["caja_min"], c["mes_caja_min"] = cm, mm
    pn, mp = pn_minimo(r36)
    c["pn363_min"], c["mes_pn363_min"] = pn, mp
    h, mh = holgura_363(r36)
    c["holgura_363"], c["mes_holgura_363"] = h, mh
    c["umbral_363"] = r36["balances"][0]["umbral_363"]
    c["van5"], c["flujos5"] = van(r60)
    c["caja_m0"] = r36["caja"][0]
    c["pn_m0"] = r36["balances"][0]["total_pn"]
    c["importe_enisa"] = r36["S"]["enisa"]["importe"]
    c["margen_fp_enisa"] = c["pn_m0"] - c["importe_enisa"]
    c["aportacion_fundadores"] = r36["S"]["capital_fundadores"] + r36["S"]["prima_fundadores"]
    c["capital"] = r36["S"]["capital_fundadores"]
    c["mes_enisa"] = r36["S"]["enisa"]["mes_cobro"]
    c["licencia"] = r36["S"]["precio_licencia"]
    c["sueldo_renunciado"] = sueldo_renunciado(r36["S"])
    return c


def cifras_texto(res, escen, extra):
    """Cifras sueltas que cita el texto de 3.9 (cifras_v11.json)."""
    S, B, P = res["S"], res["balances"], res["P"]
    y = {a: pyg_anual(res, a) for a in (1, 2, 3)}
    t = {a: tesoreria_anual(res, a) for a in (1, 2, 3)}
    c = {}
    for a in (1, 2, 3):
        for k in ("ingresos", "impl", "lic", "gastos_expl", "ebitda", "resultado", "bai", "is", "amort", "financieros"):
            c[f"{k}_a{a}"] = y[a][k]
        c[f"cash_flow_a{a}"] = y[a]["resultado"] + y[a]["amort"]
        c[f"caja_fin_a{a}"] = t[a]["final"]
        c[f"flujo_expl_a{a}"] = t[a]["explotacion"]
        c[f"activos_a{a}"] = activos_al_final(res, a * 12)
        c[f"anticipados_a{a}"] = B[a * 12]["periodificacion"] + B[a * 12]["anticipos"]
        c[f"pn_a{a}"] = B[a * 12]["total_pn"]
    c["caja_m0"] = B[0]["tesoreria"]
    c["pn_m0"] = B[0]["total_pn"]
    c["salidas_mes2"] = res["caja"][1] - res["caja"][2]
    c["salidas_mes4"] = res["caja"][3] - res["caja"][4]
    c["iva_mes7"] = P["pago_iva"][7]
    c["iva_mes10"] = P["pago_iva"][10]
    c["bin_final"] = res["bin_final"]
    c["perdidas_a1_a2"] = -(y[1]["resultado"] + y[2]["resultado"])
    c["licencias_cartera_a3"] = activos_al_final(res, 36) * S["precio_licencia"]
    pm, lic_eq = punto_muerto(res, 3)
    c["punto_muerto_licencias"] = pm
    c["clientes_equilibrio"] = math.ceil(pm)
    c["margen_a3"] = y[3]["resultado"] / y[3]["ingresos"]
    c["peso_impl_a3"] = y[3]["impl"] / y[3]["ingresos"]
    c["personal_a1_peso"] = (y[1]["sueldos"] + y[1]["cargas"]) / y[1]["gastos_expl"]
    cm, mm = caja_minima(res)
    c["caja_min"], c["mes_caja_min"] = cm, mm
    pn, mp = pn_minimo(res)
    c["pn363_min"], c["mes_pn363_min"] = pn, mp
    c["pn_mes"] = {m: B[m]["total_pn"] for m in range(0, 13)}
    c["pn363_mes"] = {m: B[m]["pn_mercantil"] for m in range(0, 13)}
    c["caja_mes"] = {m: res["caja"][m] for m in range(0, 13)}
    c["multiplo_licencias"] = 1000000.0 / c["licencias_cartera_a3"]
    c["multiplo_ingresos"] = 1000000.0 / y[3]["ingresos"]
    for n in (20, 30, 45):
        veces = n * S["precio_licencia"] * MULTIPLO_VENTA * S["ronda"]["participacion"] / S["ronda"]["importe"]
        c[f"inversor_veces_{n}"] = veces
        c[f"inversor_tir_{n}"] = veces ** (1 / ANIOS_HASTA_VENTA) - 1
    for nombre, r in escen:
        cmn, mmn = caja_minima(r, 1, 36)
        h, mh = holgura_363(r, 36)
        neg = [m for m in range(1, 37) if r["caja"][m] < 0]
        c["esc: " + nombre] = {
            "caja_min": cmn, "mes_caja_min": mmn, "holgura_363": h, "mes_holgura": mh,
            "pn363_min": h + r["balances"][mh]["umbral_363"], "primer_mes_negativo": neg[0] if neg else None,
            "resultado_a2": pyg_anual(r, 2)["resultado"], "resultado_a3": pyg_anual(r, 3)["resultado"],
            "ingresos_a3": pyg_anual(r, 3)["ingresos"], "caja_36": r["caja"][36],
        }
    c["cascada"] = {a: cascada(res, a) for a in (1, 2, 3)}
    c["punto_muerto"] = {a: punto_muerto_detalle(res, a) for a in (2, 3)}
    c["fondo_maniobra"] = {m: fondo_maniobra(res, m) for m in (0, 12, 24, 36)}
    peor = min(range(1, 12), key=lambda m: fondo_maniobra(res, m)["fondo_maniobra"])
    c["fondo_maniobra_peor_mes"] = {"mes": peor, "fondo_maniobra": fondo_maniobra(res, peor)["fondo_maniobra"]}
    c["valor_cliente"] = {n: valor_cliente(licencia(**REDES_CALIBRACION[n]), minuto) for n, minuto in MINUTO_TIEPI.items()}
    c["ipc_3_pct"] = valor_ipc(S)
    c["contrataciones_escenarios"] = {nombre: meses_contratacion(s) for nombre, s in lista_escenarios(S)}
    c.update(extra)
    return c


def texto_comparacion(c10, c11, extra):
    lin = ["Antes (v10) y después (v11) de las cifras clave del plan financiero, en euros", ""]
    filas = [
        ("Licencia anual por cliente del plan", "licencia"),
        ("Aportación de los fundadores en la constitución", "aportacion_fundadores"),
        ("Capital social", "capital"),
        ("Sueldo de los fundadores al que renuncian (año 1)", "sueldo_renunciado"),
        ("Préstamo de ENISA", "importe_enisa"),
        ("Mes de cobro de ENISA (0 = momento 0)", "mes_enisa"),
        ("Fondos propios del momento 0 (cuentas de 2026)", "pn_m0"),
        ("Margen de los fondos propios sobre el préstamo de ENISA", "margen_fp_enisa"),
        ("Tesorería del momento 0", "caja_m0"),
        ("Ingresos del año 1", "ingresos_a1"), ("Ingresos del año 2", "ingresos_a2"),
        ("Ingresos del año 3", "ingresos_a3"), ("Ingresos del año 4 (contraste)", "ingresos_a4"),
        ("Ingresos del año 5 (contraste)", "ingresos_a5"),
        ("Resultado del año 1", "resultado_a1"), ("Resultado del año 2", "resultado_a2"),
        ("Resultado del año 3", "resultado_a3"), ("Resultado del año 4 (contraste)", "resultado_a4"),
        ("Resultado del año 5 (contraste)", "resultado_a5"),
        ("Licencias en cartera al final del año 3", "licencias_cartera_a3"),
        ("Punto muerto: euros de licencias que pagan la estructura del año 3", "punto_muerto_euros"),
        ("Tesorería al final del año 3", "caja_fin_a3"),
        ("Tesorería al final del año 5 (contraste)", "caja_fin_a5"),
        ("VAN de los flujos libres de los años 1 a 5 al 10 %, sin valor residual", "van5"),
        ("Patrimonio neto al final del año 3", "pn_a3"),
        ("Umbral del art. 363 LSC (mitad del capital)", "umbral_363"),
    ]
    lin.append("| Cifra | v10 | v11 |")
    lin.append("| --- | --- | --- |")
    for nombre, k in filas:
        lin.append(f"| {nombre} | {fmt(c10[k])} | {fmt(c11[k])} |")
    lin.append(f"| Punto muerto en licencias | {fmt1(c10['punto_muerto_licencias'])} | "
               f"{fmt1(c11['punto_muerto_licencias'])} |")
    lin.append(f"| Caja más baja (mes) | {fmt(c10['caja_min'])} (mes {c10['mes_caja_min']}) | "
               f"{fmt(c11['caja_min'])} (mes {c11['mes_caja_min']}) |")
    lin.append(f"| Patrimonio neto más bajo con el participativo (mes) | {fmt(c10['pn363_min'])} (mes "
               f"{c10['mes_pn363_min']}) | {fmt(c11['pn363_min'])} (mes {c11['mes_pn363_min']}) |")
    lin.append(f"| Margen sobre la mitad del capital en el peor mes | {fmt(c10['holgura_363'])} (mes "
               f"{c10['mes_holgura_363']}) | {fmt(c11['holgura_363'])} (mes {c11['mes_holgura_363']}) |")
    lin.append("")
    lin.append(f"Flujos libres de los años 1 a 5, v10: {', '.join(fmt(f) for f in c10['flujos5'])}")
    lin.append(f"Flujos libres de los años 1 a 5, v11: {', '.join(fmt(f) for f in c11['flujos5'])}")
    lin.append("")
    lin.append(f"Con 40.000 euros y ENISA en el mes {MES_ENISA_V11}, los fundadores tienen que renunciar como mínimo a un "
               f"{fmt_pct(extra['rebaja_minima'], 1)} de su sueldo en los meses 1 a {MES_ENISA_V11 - 1}; el plan pone "
               f"un {fmt_pct(RECORTE_FUNDADORES_V11[2], 0)}.")
    lin.append(f"Importe redondo más alto de ENISA que cumple la condición de fondos propios y deja caja y patrimonio en "
               f"positivo todos los meses: {fmt(extra['enisa_maximo'])} euros; el plan pide "
               f"{fmt(IMPORTE_ENISA_V11)}.")
    lin.append("")
    lin.append("Tarifa: " + texto_tarifa(moneda="euros"))
    for nombre, red in (list(REDES_CALIBRACION.items()) + [("Red tipo del plan", RED_TIPO_PLAN)]
                        + list(REDES_REFERENCIA.items())):
        lin.append(f"  {nombre}: {fmt(red['ct'])} CT, {fmt(red['senales'])} señales -> "
                   f"{fmt(licencia(**red))} euros al año")
    if SENALES_PROVISIONALES:
        lin.append("  Las señales de las tres distribuidoras son provisionales (78.301 tags de TedisNet repartidos "
                   "por CT).")
    return "\n".join(lin) + "\n"


def mensual_csv(res, ruta):
    P, B = res["P"], res["balances"]
    with open(ruta, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["mes", "ingresos", "gastos_explotacion", "resultado_mes", "cobros", "pagos",
                    "entrada_prestamo", "entrada_capital", "tesoreria_final", "patrimonio_neto",
                    "patrimonio_neto_363", "umbral_363"])
        for m in range(0, res["N"] + 1):
            if m == 0:
                w.writerow([0, 0, 0, 0, 0, 0, 0, 0, round(res["caja"][0], 2), round(B[0]["total_pn"], 2),
                            round(B[0]["pn_mercantil"], 2), round(B[0]["umbral_363"], 2)])
                continue
            ing = P["ing_impl"][m] + P["ing_lic"][m]
            gas = P["sueldos"][m] + P["cargas"][m] + P["externos"][m] + P["comision"][m] + P["imprevistos"][m]
            resm = ing - gas - P["amort"][m] - P["financieros"][m] - P["is"][m]
            pagos = (P["pago_nomina"][m] + P["pago_ss"][m] + P["pago_irpf"][m] + P["pago_externos"][m]
                     + P["pago_iva"][m] + P["pago_is"][m] + P["pago_inversion"][m] + P["pago_intereses"][m]
                     + P["pago_principal"][m] + P["pago_gastos_ronda"][m])
            w.writerow([m, round(ing, 2), round(gas, 2), round(resm, 2), round(P["cobros"][m], 2), round(pagos, 2),
                        round(P["entrada_prestamo"][m], 2), round(P["entrada_capital"][m], 2),
                        round(res["caja"][m], 2), round(B[m]["total_pn"], 2), round(B[m]["pn_mercantil"], 2),
                        round(B[m]["umbral_363"], 2)])


def escribir_xlsx(res, ruta):
    try:
        import openpyxl
    except ImportError:
        return False
    P = res["P"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mensual"
    cols = list(P.keys())
    ws.append(["mes"] + cols + ["tesoreria_final"])
    for m in range(1, res["N"] + 1):
        ws.append([m] + [round(P[k][m], 2) for k in cols] + [round(res["caja"][m], 2)])
    ws2 = wb.create_sheet("Balances")
    claves = ["intangible", "material", "fianza", "clientes", "hp_deudora_iva", "tesoreria", "total_activo",
              "capital", "prima", "reservas", "anteriores", "resultado", "total_pn", "pn_mercantil", "umbral_363",
              "enisa_lp", "banco_lp", "deuda_cp", "anticipos", "periodificacion", "hp_acreedora", "ss_acreedora",
              "total_pn_pasivo"]
    ws2.append(["mes"] + claves)
    for m in range(0, res["N"] + 1):
        ws2.append([m] + [round(res["balances"][m][k], 2) for k in claves])
    ws3 = wb.create_sheet("Supuestos")
    for k, v in res["S"].items():
        ws3.append([k, json.dumps(v, ensure_ascii=False)])
    wb.save(ruta)
    return True


# ---------------------------------------------------------------------------
# Cálculo y comprobaciones
# ---------------------------------------------------------------------------
CLAVES_IGUAL_O_MEJOR = ["ingresos_a1", "ingresos_a2", "ingresos_a3", "ingresos_a4", "ingresos_a5",
                        "resultado_a1", "resultado_a2", "resultado_a3", "resultado_a4", "resultado_a5",
                        "van5", "caja_fin_a3", "caja_fin_a5"]


def comprobar_v11(r36, r60, S):
    """La v11 cumple todos los meses: caja en positivo, patrimonio por encima de la mitad del capital
    (art. 363.1.e LSC, con el participativo) y ENISA no pasa de los fondos propios de las cuentas de 2026."""
    for r in (r36, r60):
        h, mh = holgura_363(r)
        assert h >= 0, f"v11: patrimonio por debajo de la mitad del capital en el mes {mh} ({h:.2f})"
        cm, mm = caja_minima(r)
        assert cm >= 0, f"v11: caja negativa en el mes {mm} ({cm:.2f})"
    assert r36["balances"][0]["total_pn"] >= S["enisa"]["importe"], "ENISA por encima de los fondos propios de 2026"
    assert sum(r36["P"]["entrada_prestamo"][:S["enisa"]["mes_cobro"]]) == 0, "dinero de ENISA antes del mes de cobro"


def comprobar_igual_o_mejor(c10, c11):
    for k in CLAVES_IGUAL_O_MEJOR:
        assert c11[k] >= c10[k] - 0.01, f"la v11 empeora {k}: {c11[k]:.0f} frente a {c10[k]:.0f}"
    assert c11["punto_muerto_licencias"] <= c10["punto_muerto_licencias"] + 1e-9, \
        "la v11 necesita más clientes para el punto muerto"


def calcular():
    """Todo el cálculo, con sus comprobaciones. Lo usan main() y las pruebas."""
    r10 = simular(V10)
    r11 = simular(V11)
    r10_60 = simular(a_60_meses(V10))
    r11_60 = simular(a_60_meses(V11))
    comprobar_v11(r11, r11_60, V11)
    c10 = cifras_clave(r10, r10_60)
    c11 = cifras_clave(r11, r11_60)
    comprobar_igual_o_mejor(c10, c11)
    rebaja = rebaja_minima(V11, MES_ENISA_V11 - 1)
    assert rebaja is not None and RECORTE_FUNDADORES_V11[2] >= rebaja, "el medio sueldo ya no basta"
    sin_rebaja = copy.deepcopy(V11)
    sin_rebaja["recortes_fundadores"] = []
    r_sin = simular(sin_rebaja)
    extra = {
        "rebaja_minima": rebaja,
        "enisa_maximo": enisa_maximo(V11),
        "pn_sin_rebaja_mes": {m: r_sin["balances"][m]["pn_mercantil"] for m in range(1, MES_ENISA_V11)},
        "caja_sin_rebaja_min": caja_minima(r_sin)[0],
        "mes_caja_sin_rebaja_min": caja_minima(r_sin)[1],
    }
    assert extra["enisa_maximo"] is not None and extra["enisa_maximo"] >= IMPORTE_ENISA_V11
    return {"r10": r10, "r11": r11, "r10_60": r10_60, "r11_60": r11_60, "c10": c10, "c11": c11, "extra": extra}


def generar(R=None):
    """Las salidas de texto (tablas y comparación) tal como las escribe main(). Las pruebas las
    comparan con los ficheros de salida/ para saber si están al día con el código."""
    R = R or calcular()
    e11 = escenarios(V11)
    e10 = escenarios(V10)
    t11, _, _ = construir_tablas(R["r11"], e11)
    t10, _, _ = construir_tablas(R["r10"], e10)
    textos = {
        "tablas_v11.txt": "\n\n".join(t11.values()) + "\n",
        "tablas_v10.txt": "\n\n".join(t10.values()) + "\n",
        "comparacion.txt": texto_comparacion(R["c10"], R["c11"], R["extra"]),
    }
    return R, e11, t11, textos


def main():
    R, e11, t11, textos = generar()
    comp = textos["comparacion.txt"]
    if "--sin-salida" in sys.argv:
        print(comp)
        for v in t11.values():
            print()
            print(v)
        return
    os.makedirs(SALIDA, exist_ok=True)
    for nombre, texto in textos.items():
        with open(os.path.join(SALIDA, nombre), "w", encoding="utf-8", newline="\n") as f:
            f.write(texto)

    def redondear(v):
        if isinstance(v, float):
            return round(v, 4)
        if isinstance(v, dict):
            return {str(k): redondear(x) for k, x in v.items()}
        if isinstance(v, list):
            return [redondear(x) for x in v]
        return v
    cif = cifras_texto(R["r11"], e11, R["extra"])
    cif["clave"] = {k: v for k, v in R["c11"].items() if k != "flujos5"}
    cif["clave_v10"] = {k: v for k, v in R["c10"].items() if k != "flujos5"}
    with open(os.path.join(SALIDA, "cifras_v11.json"), "w", encoding="utf-8") as f:
        json.dump(redondear(cif), f, ensure_ascii=False, indent=1, sort_keys=True)
    mensual_csv(R["r11"], os.path.join(SALIDA, "mensual_v11.csv"))
    escribir_xlsx(R["r11"], os.path.join(SALIDA, "detalle_v11.xlsx"))
    print(comp)
    print("Comprobaciones superadas: balances cuadrados en los 37 cierres (y en los 61 de la extensión), "
          "tesorería y resultados encadenados; en la v11, caja y patrimonio (art. 363 LSC) en positivo todos los "
          "meses, ENISA por debajo de los fondos propios de 2026 y el plan igual o mejor que el de la v10.")


if __name__ == "__main__":
    main()
