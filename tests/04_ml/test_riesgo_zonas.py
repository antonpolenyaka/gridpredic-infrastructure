"""
Test of the risk below the distribuidora (riesgo_zonas): the zone table and
its history, the rescaling that keeps the probability of the day, the share
of each CT and the risk of the municipalities, the expected impact, the
window of the day, the mix of causes, the cascade rule and the bases of the
crews. Pure pandas, no Spark.

    python -m pytest tests/04_ml/test_riesgo_zonas.py -q
"""

import os
import sys

import numpy as np
import pandas as pd
import pytest

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
ML_DIR = os.path.join(ROOT, "etl", "jobs", "04_ml")
if ML_DIR not in sys.path:
    sys.path.insert(0, ML_DIR)

import riesgo_zonas as rz  # noqa: E402


@pytest.fixture
def network():
    """
    Six CTs of one distribuidora in two lines: line 10 (CTs a, b, c in
    Montánchez and Torre) and line 11 (CTs d, e, f in Zafra), and the
    cells of their municipalities.
    """
    dim = pd.DataFrame({
        "distribuidora_id": [1] * 6,
        "ct_id": ["a", "b", "c", "d", "e", "f"],
        "grupo_red_id": [10.0, 10.0, 10.0, 11.0, 11.0, 11.0],
        "municipio_nombre": ["Montánchez", "Montánchez", "Torre", "Zafra", "Zafra", "Zafra"],
        "potencia_kva": [100.0, 250.0, 50.0, 400.0, 160.0, 630.0],
        "num_abonados": [40, 90, 10, 200, 60, 300],
        "tipo_zona_codigo": [4, 4, 4, 3, 3, 3],
    })
    cells = pd.DataFrame({
        "distribuidora_id": [1] * 6,
        "ct_id": ["a", "b", "c", "d", "e", "f"],
        "celda_id": ["c1", "c1", "c1", "c2", "c2", "c2"],
        "ct_lat_municipio": [39.3, 39.3, 39.4, 38.4, 38.4, 38.4],
        "ct_lon_municipio": [-6.1, -6.1, -6.1, -6.4, -6.4, -6.4],
        "codigo_ine": ["10128", "10128", "10190", "06158", "06158", "06158"],
    })
    return dim, cells


def events_of(rows):
    """events as the job passes them: one row per CT cut by an incident."""
    return pd.DataFrame(rows, columns=["distribuidora_id", "periodo_id", "incidencia_id", "ct_id", "inicio_ts",
                                       "conocido_ts", "causa", "duracion_winsor_s"])


@pytest.fixture
def events():
    t = pd.Timestamp
    return events_of([
        # Line 10: three incidents of a local cause, the first cuts every CT, the others only a.
        (1, 1, 1, "a", t("2025-01-10 14:00"), t("2025-01-10 14:05"), "meteorologia", 3600),
        (1, 1, 1, "b", t("2025-01-10 14:00"), t("2025-01-10 14:05"), "meteorologia", 3600),
        (1, 1, 1, "c", t("2025-01-10 14:00"), t("2025-01-10 14:05"), "meteorologia", 3600),
        (1, 1, 2, "a", t("2025-01-10 15:30"), t("2025-01-10 15:35"), "equipo", 7200),
        (1, 2, 3, "a", t("2025-03-01 09:00"), t("2025-03-01 09:10"), "fauna", 1800),
        # Line 11: one incident of weather and one upstream (not local, left out).
        (1, 3, 4, "d", t("2025-02-20 03:00"), t("2025-02-20 03:02"), "vegetacion", 5400),
        (1, 3, 5, "e", t("2025-02-21 03:00"), t("2025-02-21 03:02"), "aguas_arriba", 600),
    ])


def test_the_zone_table_names_the_lines_by_their_municipalities(network):
    dim, cells = network
    ct = rz.zone_table(dim, cells)
    assert ct["zona"].tolist() == ["10", "10", "10", "11", "11", "11"]
    assert ct["rural"].tolist() == [1.0, 1.0, 1.0, 0.0, 0.0, 0.0]
    zones = rz.zones_of(ct)
    assert zones["z_n_ct"].tolist() == [3, 3]
    assert zones.loc[0, "nombre"] == "Montánchez / Torre"
    assert zones.loc[1, "nombre"] == "Zafra"
    assert zones.loc[0, "municipios"] == ["Montánchez", "Torre"]
    assert zones.loc[0, "z_abonados"] == 140 and zones.loc[1, "z_kva"] == 1190

    # A CT without zone or without cell is left out.
    dim2 = pd.concat([dim, pd.DataFrame([{"distribuidora_id": 1, "ct_id": "g", "grupo_red_id": np.nan,
                                           "municipio_nombre": "Zafra", "potencia_kva": 1.0, "num_abonados": 1,
                                           "tipo_zona_codigo": 3}])])
    assert len(rz.zone_table(dim2, cells)) == 6


def test_local_starts_and_history_count_what_was_known_before_the_day(network, events):
    dim, cells = network
    ct = rz.zone_table(dim, cells)
    starts = rz.local_starts(events, ct)
    # One row per incident and zone; the upstream one is not local.
    assert len(starts) == 4
    assert sorted(starts["incidencia_id"].tolist()) == [1, 2, 3, 4]
    assert starts["fecha"].min() == pd.Timestamp("2025-01-10")

    zones = rz.zones_of(ct)
    grid = zones[["distribuidora_id", "zona"] + rz.ZONE_FEATURES_SIZE].merge(
        pd.DataFrame({"fecha": pd.to_datetime(["2025-01-10", "2025-01-11", "2025-03-02"])}), how="cross")
    hist = rz.zone_history(grid, starts)
    row = hist.set_index(["zona", "fecha"])
    # The two incidents of 10.01 are known at 14:05 and 15:35, after the start of the day: not counted that day.
    assert row.loc[("10", pd.Timestamp("2025-01-10")), "z_hist_30d"] == 0
    assert row.loc[("10", pd.Timestamp("2025-01-11")), "z_hist_30d"] == 2
    assert row.loc[("10", pd.Timestamp("2025-03-02")), "z_hist_30d"] == 1
    assert row.loc[("10", pd.Timestamp("2025-03-02")), "z_hist_365d"] == 3
    assert row.loc[("10", pd.Timestamp("2025-03-02")), "z_hist_dias_desde"] == pytest.approx(1 - 9.0 / 24 - 10.0 / 1440, abs=0.01)
    assert row.loc[("11", pd.Timestamp("2025-01-10")), "z_hist_dias_desde"] == 3650


def test_zone_days_builds_the_grid_with_weather_and_labels(network, events):
    dim, cells = network
    ct = rz.zone_table(dim, cells)
    zones = rz.zones_of(ct)
    starts = rz.local_starts(events, ct)
    cell_days = pd.DataFrame({
        "celda_id": ["c1", "c2", "c1", "c2"],
        "fecha": pd.to_datetime(["2025-01-10", "2025-01-10", "2025-01-11", "2025-01-11"]),
        "racha_max": [80.0, 30.0, 20.0, np.nan], "clima_max": [2.5, 1.0, 0.5, np.nan],
        "lluvia": [4.0, 0.0, 0.0, np.nan], "temp_max": [12.0, 15.0, 9.0, np.nan],
    })
    weather = rz.zone_weather(ct, cell_days)
    assert set(weather.columns) >= set(rz.ZONE_FEATURES_WEATHER)
    w = weather.set_index(["zona", "fecha"])
    assert w.loc[("10", pd.Timestamp("2025-01-10")), "z_racha_max"] == 80.0
    assert w.loc[("10", pd.Timestamp("2025-01-10")), "z_racha_media"] == 80.0
    # A cell without weather gives no mean (not zero).
    assert np.isnan(w.loc[("11", pd.Timestamp("2025-01-11")), "z_racha_media"])

    days = pd.DataFrame({"distribuidora_id": [1, 1], "fecha": pd.to_datetime(["2025-01-10", "2025-01-11"]),
                         "dia_racha_max": [80.0, 20.0]})
    table = rz.zone_days(zones, starts, weather, days, ["dia_racha_max"], "2025-01-10", "2025-01-11")
    assert len(table) == 4
    assert table.set_index(["zona", "fecha"])["y"].to_dict() == {
        ("10", pd.Timestamp("2025-01-10")): 1, ("10", pd.Timestamp("2025-01-11")): 0,
        ("11", pd.Timestamp("2025-01-10")): 0, ("11", pd.Timestamp("2025-01-11")): 0}
    assert table["dist_1"].eq(1.0).all()
    feats = rz.zone_features(zones, ["dia_racha_max"])
    assert feats[0] == "dist_1" and feats[-1] == "dia_racha_max"
    assert all(f in table.columns for f in feats)


def test_rescale_keeps_the_probability_of_the_day_and_the_order():
    scores = np.array([0.02, 0.10, 0.05])
    for p_day in (0.05, 0.3, 0.9):
        p = rz.rescale_to_day(scores, p_day)
        assert 1 - np.prod(1 - p) == pytest.approx(p_day, abs=1e-6)
        assert np.argsort(p).tolist() == np.argsort(scores).tolist()
    # Out of range or empty: nothing moves.
    assert rz.rescale_to_day(scores, 0.0).tolist() == pytest.approx(scores.tolist())
    assert len(rz.rescale_to_day([], 0.5)) == 0


def test_levels_follow_the_thresholds():
    levels = {"muy alto": 0.95, "alto": 0.9, "medio": 0.75}
    assert [rz.level_of(p, levels) for p in (0.2, 0.75, 0.9, 0.99)] == ["bajo", "medio", "alto", "muy alto"]


def test_ct_shares_and_municipality_risk(network, events):
    dim, cells = network
    ct = rz.zone_table(dim, cells)
    shares = rz.ct_shares(ct, events, prior=0.5, weight=5.0)
    # Line 10: 3 incidents, a in all 3, b and c in 1; the mean share of the line is (3+1+1)/3/3 = 5/9.
    prior = 5 / 9
    assert shares[(1, "a")] == pytest.approx((3 + 5 * prior) / (3 + 5))
    assert shares[(1, "b")] == pytest.approx((1 + 5 * prior) / (3 + 5))
    assert shares[(1, "a")] > shares[(1, "b")] == shares[(1, "c")]
    # Line 11: one local incident, on d; e and f never cut (the upstream one does not count).
    assert shares[(1, "d")] > shares[(1, "e")] == shares[(1, "f")]
    assert 0.02 <= shares.min() and shares.max() <= 1.0

    zone_p = pd.DataFrame({"distribuidora_id": [1, 1], "zona": ["10", "11"], "probabilidad": [0.2, 0.1]})
    muni = rz.municipality_risk(ct, zone_p, shares)
    assert muni["municipio"].tolist() == ["Montánchez", "Torre", "Zafra"]
    m = muni.set_index("municipio")
    # Montánchez has a and b of the same line: the probability of the line times the largest share (a).
    assert m.loc["Montánchez", "probabilidad"] == pytest.approx(0.2 * shares[(1, "a")])
    assert m.loc["Torre", "probabilidad"] == pytest.approx(0.2 * shares[(1, "c")])
    assert m.loc["Zafra", "n_ct"] == 3 and m.loc["Zafra", "zonas"] == ["11"]
    assert m.loc["Montánchez", "clientes_esperados"] == pytest.approx(0.2 * (shares[(1, "a")] * 40 + shares[(1, "b")] * 90))
    assert m.loc["Montánchez", "codigo_ine"] == "10128"

    cts = rz.ct_risk(ct, zone_p, shares)
    assert cts.set_index("ct_id").loc["a", "probabilidad"] == pytest.approx(0.2 * shares[(1, "a")])
    assert (cts["probabilidad"] <= cts["zona"].map({"10": 0.2, "11": 0.1})).all()


def test_impact_takes_the_medians_of_the_distribuidora_for_the_zones_with_few_incidents(network, events):
    dim, cells = network
    ct = rz.zone_table(dim, cells)
    stats = rz.impact_stats(ct, events)
    s = stats.set_index("zona")
    assert s.loc["10", "n"] == 3 and s.loc["10", "cts"] == 1 and s.loc["10", "horas"] == 1.0
    # Line 11 has one incident: it takes the medians of the distribuidora (over its 4 incidents).
    assert s.loc["11", "n"] == 1
    assert s.loc["11", "cts"] == 1 and s.loc["11", "horas"] == pytest.approx(1.25)
    zone_p = pd.DataFrame({"distribuidora_id": [1, 1], "zona": ["10", "11"], "probabilidad": [0.5, 0.2]})
    out = rz.impact(zone_p, stats, installed_kva=1590.0)
    assert out["incidencias_esperadas"] == pytest.approx(0.7)
    assert out["ct_esperados"] == pytest.approx(0.5 * 1 + 0.2 * 1)
    assert out["minutos_tiepi_esperados"] == pytest.approx(60 * out["kva_h_esperados"] / 1590.0)
    # With the stats already merged the result is the same.
    merged = zone_p.merge(stats, on=["distribuidora_id", "zona"])
    assert rz.impact(merged, stats, 1590.0) == out
    assert rz.impact(zone_p, stats, 0.0)["minutos_tiepi_esperados"] is None


def test_the_window_of_the_day_follows_the_gust_peak():
    profile = pd.DataFrame({"hora": range(24), "racha_max": [20.0] * 10 + [50, 70, 85, 80, 60, 40] + [25.0] * 8})
    w = rz.risk_window(profile, share=0.8, calm_kmh=40)
    # 80 % of 85 = 68: hours 11 - 13, widened to 10 - 14.
    assert (w["inicio"], w["fin"], w["pico"], w["racha_pico"]) == (10, 14, 12, 85.0)
    assert w["texto"].startswith("de 10:00 a 15:00")
    calm = rz.risk_window(profile.assign(racha_max=profile["racha_max"] / 3), calm_kmh=40)
    assert calm["inicio"] is None and "poco viento" in calm["texto"]
    assert rz.risk_window(pd.DataFrame(columns=["hora", "racha_max"]))["texto"] == "sin perfil horario"
    edge = rz.risk_window(pd.DataFrame({"hora": range(24), "racha_max": [90.0] + [10.0] * 23}))
    assert (edge["inicio"], edge["fin"]) == (0, 1)


def test_lags_cause_mix_and_gust_buckets(network, events):
    dim, cells = network
    ct = rz.zone_table(dim, cells)
    starts = rz.local_starts(events, ct)
    cell_days = pd.DataFrame({
        "celda_id": ["c1", "c1", "c2"],
        "fecha": pd.to_datetime(["2025-01-10", "2025-03-01", "2025-02-20"]),
        "racha_max": [75.0, 30.0, 50.0],
        "hora_pico": pd.to_datetime(["2025-01-10 13:00", "2025-03-01 18:00", "2025-02-20 07:00"]),
    })
    lags = rz.lag_stats(starts, cell_days, ct, windy_kmh=60)
    # The weather incidents start 1 h after the peak (10.01) and 4 h before it (20.02).
    assert lags["meteorologia"]["n"] == 2 and lags["meteorologia"]["en_3h"] == 0.5
    assert lags["meteorologia"]["en_6h"] == 1.0 and lags["meteorologia"]["mediana_h"] == pytest.approx(-1.5)
    assert lags["dias_ventosos"]["n"] == 2 and lags["dias_ventosos"]["en_3h"] == 1.0
    assert lags["todas"]["n"] == 4
    assert rz.lag_stats(starts, cell_days.drop(columns="hora_pico"), ct) == {}

    mix = rz.cause_mix(starts, ct, cell_days)
    strong = mix[mix["bucket"] == "70 km/h o más"].set_index("familia")
    assert strong.loc["meteorologia", "n"] == 1 and strong.loc["equipo", "n"] == 1
    assert strong["parte"].sum() == pytest.approx(1.0)
    # vegetacion belongs to the weather family.
    mid = mix[mix["bucket"] == "40 - 55 km/h"].set_index("familia")
    assert mid.loc["meteorologia", "parte"] == 1.0
    assert [rz.gust_bucket(g) for g in (10, 40, 54.9, 69, 70, 120, float("nan"))] == [
        "menos de 40 km/h", "40 - 55 km/h", "40 - 55 km/h", "55 - 70 km/h", "70 km/h o más", "70 km/h o más",
        "menos de 40 km/h"]


def test_the_cascade_rule_counts_the_incidents_followed_by_another_one(network, events):
    dim, cells = network
    ct = rz.zone_table(dim, cells)
    starts = rz.local_starts(events, ct)
    rule = rz.cascade_rule(starts, hours=3)
    # Of the four incidents only the first of 10.01 (14:00) has another one within 3 h (15:30).
    assert rule["n"] == 4 and rule["todas"] == 0.25
    assert rule["meteorologia"] == 0.5  # the one of 10.01 yes, the one of 20.02 no
    assert rz.cascade_rule(starts.iloc[0:0]) == {"horas": 3.0}


def test_crew_plan_puts_the_bases_where_the_risk_is():
    municipalities = pd.DataFrame({
        "municipio": ["Norte", "Centro", "Sur", "Lejos"],
        "lat": [39.5, 39.0, 38.5, 40.5], "lon": [-6.0, -6.0, -6.0, -6.0],
        "n_ct": [10, 20, 10, 2], "probabilidad": [0.1, 0.3, 0.1, 0.01],
        "clientes_esperados": [50.0, 400.0, 50.0, 1.0],
    })
    one = rz.crew_plan(municipalities, 1, speed_kmh=60, road_factor=1.3)
    assert one["n"] == 1 and one["bases"][0]["municipio"] == "Centro"
    assert one["bases"][0]["cubre_municipios"] == 4 and one["bases"][0]["parte_riesgo"] == 1.0
    assert one["llegada_max_min"] == pytest.approx(rz.haversine_km(np.array([[39.0, -6.0]]), np.array([[40.5, -6.0]]))[0, 0] * 1.3)
    two = rz.crew_plan(municipalities, 2, fixed=["Lejos"])
    assert two["n"] == 2 and two["llegada_media_min"] < one["llegada_media_min"]
    assert sum(b["parte_riesgo"] for b in two["bases"]) == pytest.approx(1.0)
    assert sum(b["cubre_municipios"] for b in two["bases"]) == 4
    assert two["fijas"]["municipios"] == ["Lejos"]
    assert two["fijas"]["llegada_media_min"] > two["llegada_media_min"]
    assert 0 <= two["riesgo_30min"] <= two["riesgo_45min"] <= 1
    assert two["bases"][0]["municipios"][0] == two["bases"][0]["municipio"] or len(two["bases"][0]["municipios"]) >= 1
    assert sum(len(b["municipios"]) for b in two["bases"]) == 4
    curve = rz.crew_curve(municipalities, 3)
    assert [c["n"] for c in curve] == [1, 2, 3] and curve[0]["bases"] == ["Centro"]
    assert curve[0]["llegada_media_min"] >= curve[1]["llegada_media_min"] >= curve[2]["llegada_media_min"]
    assert rz.positioning({"inicio": 10, "fin": 14}, two).startswith("brigadas en sus bases a las 9:00")
    assert rz.positioning({"inicio": 0, "fin": 3}, two).startswith("brigadas en sus bases a las 0:00")
    assert rz.positioning({"inicio": None}, two).startswith("sin franja")
    assert rz.positioning({"inicio": 10}, {}) == ""
    # With the minutes by road the plan says so and uses them; a table that misses many municipalities is ignored.
    names = municipalities["municipio"].tolist()
    road = pd.DataFrame(np.full((4, 4), 90.0), index=names, columns=names)
    np.fill_diagonal(road.to_numpy(), 0.0)
    road.loc["Norte", "Centro"] = road.loc["Centro", "Norte"] = 5.0
    by_road = rz.crew_plan(municipalities, 1, road_minutes=road)
    assert by_road["tiempos"].startswith("carretera") and one["tiempos"].startswith("línea recta")
    assert by_road["bases"][0]["municipio"] == "Centro" and by_road["llegada_max_min"] == 90.0
    partial = rz.crew_plan(municipalities, 1, road_minutes=road.iloc[:2, :2])
    assert partial["tiempos"] == one["tiempos"] and partial["llegada_media_min"] == one["llegada_media_min"]
    # Without weights the CTs decide; without coordinates or crews nothing.
    assert rz.crew_plan(municipalities.assign(clientes_esperados=0.0), 1)["bases"][0]["municipio"] == "Centro"
    assert rz.crew_plan(municipalities, 0) == {}
    assert rz.crew_plan(municipalities.assign(lat=np.nan), 1) == {}


def test_the_shift_sheet_has_one_sheet_per_table(tmp_path):
    payload = {
        "fecha": "2026-01-15", "fecha_texto": "15.01.2026", "generado": "14.01.2026 19:40 UTC", "run_id": "manana_x",
        "reglas": {"texto": "Tras una avería, otra en 3 h en el 11 % de los casos."}, "notas": ["Calser hasta el 09.08.2026."],
        "distribuidoras": [{
            "id": 3, "nombre": "EOSA", "nivel": "alto", "probabilidad": 0.31, "veces_tasa_base": 4.1, "esperadas": 0.4,
            "motivos": "racha prevista (x3)", "impacto": {"ct_esperados": 1.2, "clientes_esperados": 80.0, "kva_esperados": 300.0},
            "franja": {"texto": "de 10:00 a 15:00", "pico": 12, "racha_pico": 85.0},
            "causas": [{"familia": "meteorología y vegetación", "parte": 0.6}, {"familia": "equipos", "parte": 0.4}],
            "brigadas": {"n": 2, "tiempos": "carretera (OSRM sobre OpenStreetMap)", "llegada_media_min": 16.0,
                         "posicionamiento": "brigadas en sus bases a las 9:00",
                         "bases": [{"municipio": "Montánchez", "cubre_municipios": 3, "cubre_ct": 40, "parte_riesgo": 0.6,
                                    "llegada_media_min": 12.0, "llegada_max_min": 30.0, "municipios": ["Montánchez", "Torre"]}],
                         "curva": [{"n": 1, "llegada_media_min": 40.0, "riesgo_30min": 0.5, "llegada_max_min": 90.0, "bases": ["Montánchez"]},
                                   {"n": 2, "llegada_media_min": 16.0, "riesgo_30min": 0.9, "llegada_max_min": 40.0, "bases": ["Montánchez", "Hervás"]}]},
            "zonas": [{"zona": "10", "nombre": "Montánchez / Torre", "nivel": "alto", "probabilidad": 0.2, "parte": 0.6, "n_ct": 30,
                       "abonados": 900.0, "kva": 5000.0, "clientes_esperados": 60.0, "racha_max": 85.0, "lluvia": 2.0,
                       "hist_365d": 4.0, "hist_dias_desde": 40.0, "municipios": ["Montánchez", "Torre"]},
                      {"zona": "11", "nombre": "Hervás", "nivel": "medio", "probabilidad": 0.14, "parte": 0.4, "n_ct": 20,
                       "abonados": 600.0, "kva": 3000.0, "clientes_esperados": 30.0, "racha_max": 70.0, "lluvia": 0.0,
                       "hist_365d": 1.0, "hist_dias_desde": 200.0, "municipios": ["Hervás"]}],
            "municipios": [{"nombre": "Montánchez", "codigo_ine": "10128", "nivel": "alto", "probabilidad": 0.18, "n_ct": 20,
                            "abonados": 700.0, "clientes_esperados": 50.0, "racha_max": 85.0, "zonas": ["Montánchez / Torre"]}],
            "cts": [{"ct_id": "a", "municipio": "Montánchez", "zona": "10", "abonados": 40.0, "kva": 100.0, "probabilidad": 0.15}],
        }],
    }
    sheets = rz.shift_sheets(payload)
    assert list(sheets) == ["Resumen", "Lineas", "Municipios", "CT", "Brigadas", "Brigadas por numero", "Notas"]
    assert sheets["Resumen"].loc[0, "bases"] == "Montánchez" and sheets["Resumen"].loc[0, "causas (días como mañana)"].startswith("meteorología")
    assert sheets["Lineas"]["línea"].tolist() == ["Montánchez / Torre", "Hervás"]
    assert sheets["Brigadas por numero"]["brigadas"].tolist() == [1, 2]
    assert sheets["Notas"]["nota"].iloc[0].startswith("Riesgo de mañana 15.01.2026")
    path = tmp_path / "parte_guardia.xlsx"
    assert rz.write_shift_sheet(payload, str(path)) and path.exists()
    from openpyxl import load_workbook

    book = load_workbook(str(path))
    assert book.sheetnames == list(sheets)
    assert book["Lineas"]["E2"].number_format == "0.0%" and book["Lineas"]["E2"].value == 0.2
    assert book["Resumen"].freeze_panes == "A2"
    # An empty table leaves a sheet that says so.
    empty = {**payload, "distribuidoras": []}
    assert rz.write_shift_sheet(empty, str(tmp_path / "vacio.xlsx"))
    assert load_workbook(str(tmp_path / "vacio.xlsx"))["CT"]["A1"].value == "sin datos"


def test_clean_makes_the_payload_json_friendly(tmp_path):
    payload = {"a": np.int64(3), "b": np.float32(1.5), "c": float("nan"), "d": pd.Timestamp("2026-01-15 20:30"),
               "e": [np.float64(2.0), pd.NaT], "f": {1: "x"}, "g": np.array([1, 2])}
    out = rz.clean(payload)
    assert out == {"a": 3, "b": 1.5, "c": None, "d": "2026-01-15 20:30", "e": [2.0, None], "f": {"1": "x"}, "g": [1, 2]}
    path = tmp_path / "riesgo.json"
    rz.write_json(payload, str(path))
    assert path.read_text(encoding="utf-8").startswith("{")
    assert rz.now_text().endswith("UTC")
