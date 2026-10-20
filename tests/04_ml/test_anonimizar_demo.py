"""
web/anonimizar_demo.py: the copy of the public demo leaves with pseudonyms for the towns, their INE codes and
the AEMET stations, the same in every file and in every copy, and the dashboard can still match the zones with
their towns. No Spark.
"""

import importlib.util
import json
import os
import re

from openpyxl import Workbook, load_workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SPEC = importlib.util.spec_from_file_location("anonimizar_demo", os.path.join(ROOT, "web", "anonimizar_demo.py"))
anon = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(anon)

TOWN = re.compile(r"^Municipio \d+$")


def day(names_first: bool = True) -> dict:
    towns = [{"nombre": "Villa Alta", "codigo_ine": "10001"}, {"nombre": "Alta", "codigo_ine": "10002"}]
    return {
        "distribuidoras": [{
            "id": 1, "nombre": "Distribuidora 1",
            "municipios": towns if names_first else towns[::-1],
            "zonas": [{"zona": "5", "nombre": "Villa Alta / Alta", "municipios": ["Villa Alta", "Alta"]}],
            # A CT whose id is the same number as an INE code: it must not change.
            "cts": [{"ct_id": "10001", "municipio": "Alta"}],
            "zonas_racha": "Villa Alta (60 km/h, 1,10 veces su p98); Alta (55 km/h, 0,98 veces su p98)",
            "brigadas": {"bases": [{"municipio": "Alta", "municipios": ["Alta", "Villa Alta"]}]},
        }],
    }


def write(path: str, data, indent=1) -> None:
    with open(path, "w", encoding="utf-8") as handle:
        json.dump(data, handle, ensure_ascii=False, indent=indent)
        handle.write("\n")


def folder_with_data(tmp_path) -> str:
    data = tmp_path / "data"
    data.mkdir()
    write(data / "riesgo.json", day())
    write(data / "riesgo_20260128.json", day(names_first=False))
    write(data / "observado.json", {"estaciones": [{"idema": "1234X", "nombre": "VILLA ALTA"}],
                                    "distribuidoras": [{"texto": "racha en VILLA ALTA"}]})
    write(data / "sector.json", {"ambitos": ["Alta", "España"], "mezcla": {"Alta": 0.5}})
    book = Workbook()
    sheet = book.active
    sheet.title = "Municipios"
    sheet.append(["distribuidora", "municipio", "código INE", "CT"])
    sheet.append(["Distribuidora 1", "Villa Alta", "10001", "10002"])
    book.save(data / "parte_guardia.xlsx")
    return str(data)


def snapshot(folder: str) -> dict:
    return {n: open(os.path.join(folder, n), "rb").read() for n in sorted(os.listdir(folder))}


def test_pseudonyms_everywhere(tmp_path):
    folder = folder_with_data(tmp_path)
    mapa = str(tmp_path / "mapa.json")
    sector_before = open(os.path.join(folder, "sector.json"), "rb").read()

    assert anon.main([folder, "--mapa", mapa]) == 0

    mapping = json.load(open(mapa, encoding="utf-8"))
    villa = mapping["municipios"]["Villa Alta"]["seudonimo"]
    alta = mapping["municipios"]["Alta"]["seudonimo"]
    assert TOWN.match(villa) and TOWN.match(alta) and villa != alta
    for name in ("riesgo.json", "riesgo_20260128.json"):
        dist = json.load(open(os.path.join(folder, name), encoding="utf-8"))["distribuidoras"][0]
        assert {t["nombre"] for t in dist["municipios"]} == {villa, alta}
        assert {t["codigo_ine"] for t in dist["municipios"]} == {mapping["municipios"][n]["codigo"] for n in ("Villa Alta", "Alta")}
        assert all(t["codigo_ine"].startswith("99") and len(t["codigo_ine"]) == 5 for t in dist["municipios"])
        # The zone keeps the name made of its towns, so the dashboard still finds them.
        assert dist["zonas"][0]["nombre"] == f"{villa} / {alta}" and dist["zonas"][0]["municipios"] == [villa, alta]
        assert dist["cts"][0] == {"ct_id": "10001", "municipio": alta}
        assert dist["zonas_racha"].startswith(f"{villa} (60 km/h") and f"; {alta} (55 km/h" in dist["zonas_racha"]
        assert dist["brigadas"]["bases"][0] == {"municipio": alta, "municipios": [alta, villa]}

    observed = json.load(open(os.path.join(folder, "observado.json"), encoding="utf-8"))
    station = mapping["estaciones"]["VILLA ALTA"]
    assert observed["estaciones"][0] == {"idema": station["indicativo"], "nombre": station["seudonimo"]}
    assert observed["distribuidoras"][0]["texto"] == "racha en " + station["seudonimo"]

    # The sector compares with the province, the region and Spain: untouched.
    assert open(os.path.join(folder, "sector.json"), "rb").read() == sector_before

    sheet = load_workbook(os.path.join(folder, "parte_guardia.xlsx"))["Municipios"]
    assert [c.value for c in sheet[2]] == ["Distribuidora 1", villa, mapping["municipios"]["Villa Alta"]["codigo"], "10002"]

    # Indented files stay indented, so that only the names change in the diff.
    assert open(os.path.join(folder, "riesgo.json"), encoding="utf-8").read().startswith('{\n "distribuidoras"')


def test_second_copy_keeps_the_numbers_and_changes_nothing(tmp_path):
    folder = folder_with_data(tmp_path)
    mapa = str(tmp_path / "mapa.json")
    assert anon.main([folder, "--mapa", mapa]) == 0
    first = snapshot(folder)
    numbers = json.load(open(mapa, encoding="utf-8"))

    assert anon.main([folder, "--mapa", mapa]) == 0
    assert snapshot(folder) == first

    # A new copy from the jobs, with the real names again, gets the same pseudonyms.
    write(os.path.join(folder, "riesgo_20260129.json"), day())
    assert anon.main([folder, "--mapa", mapa]) == 0
    dist = json.load(open(os.path.join(folder, "riesgo_20260129.json"), encoding="utf-8"))["distribuidoras"][0]
    assert dist["municipios"][0]["nombre"] == numbers["municipios"]["Villa Alta"]["seudonimo"]
    assert json.load(open(mapa, encoding="utf-8")) == numbers


def test_what_is_left_is_found_without_case_or_accents(tmp_path):
    folder = folder_with_data(tmp_path)
    mapa = str(tmp_path / "mapa.json")
    assert anon.main([folder, "--mapa", mapa]) == 0
    mapping = json.load(open(mapa, encoding="utf-8"))
    assert anon.leftovers(folder, mapping) == []

    write(os.path.join(folder, "nota.json"), {"texto": "avería cerca de VÍLLA alta", "codigo_ine": "10002"})
    left = anon.leftovers(folder, mapping)
    assert len(left) == 2 and all(item.startswith("nota.json") for item in left)
